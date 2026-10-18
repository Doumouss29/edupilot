"""EduPilot — plateforme de gestion scolaire multi-établissements."""
import os, sqlite3, secrets, unicodedata, re, json, datetime, calendar, functools, smtplib, csv, io, hashlib, time, hmac, base64, struct
from openpyxl import load_workbook
from flask import Response
from contextlib import contextmanager
from pathlib import Path
import requests
from urllib.parse import urlsplit, quote
from flask import Flask, request, redirect, url_for, render_template, flash, session, abort, jsonify, make_response
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from werkzeug.middleware.proxy_fix import ProxyFix
from werkzeug.exceptions import HTTPException
from flask import send_file

ROOT = Path(__file__).parent
app = Flask(__name__)
app.secret_key = os.getenv('SECRET_KEY') or secrets.token_hex(32)
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE='Lax',
    SESSION_COOKIE_SECURE=os.getenv('PUBLIC_BASE_URL','').startswith('https://'),
)
app.config['MAX_CONTENT_LENGTH'] = 12 * 1024 * 1024
DB = os.getenv('DATABASE_PATH', '/tmp/edupilot.db')
UPLOAD_DIR = Path(os.getenv('UPLOAD_DIR', '/tmp/edupilot_uploads'))
UPLOAD_DIR.mkdir(parents=True,exist_ok=True)

SCHEMA = """
PRAGMA foreign_keys=ON;
CREATE TABLE IF NOT EXISTS platform_branding(id INTEGER PRIMARY KEY CHECK(id=1),logo BLOB NOT NULL,logo_mime TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS schools(id INTEGER PRIMARY KEY, name TEXT NOT NULL, slug TEXT NOT NULL UNIQUE, plan TEXT NOT NULL DEFAULT 'standard', active INTEGER NOT NULL DEFAULT 1, subscription_end TEXT, logo BLOB, logo_mime TEXT, campus_enabled INTEGER NOT NULL DEFAULT 0);
CREATE TABLE IF NOT EXISTS campuses(id INTEGER PRIMARY KEY, school_id INTEGER NOT NULL REFERENCES schools(id), name TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY, school_id INTEGER REFERENCES schools(id), name TEXT NOT NULL, email TEXT NOT NULL, password_hash TEXT NOT NULL, role TEXT NOT NULL, matricule TEXT, active INTEGER NOT NULL DEFAULT 1, first_name TEXT, last_name TEXT, phone TEXT, must_change_password INTEGER NOT NULL DEFAULT 0, first_login_at TEXT);
CREATE TABLE IF NOT EXISTS years(id INTEGER PRIMARY KEY, school_id INTEGER NOT NULL REFERENCES schools(id), label TEXT NOT NULL, is_active INTEGER NOT NULL DEFAULT 1);
CREATE TABLE IF NOT EXISTS classes(id INTEGER PRIMARY KEY, school_id INTEGER NOT NULL REFERENCES schools(id), year_id INTEGER NOT NULL REFERENCES years(id), campus_id INTEGER REFERENCES campuses(id), name TEXT NOT NULL, cycle TEXT NOT NULL DEFAULT 'semestre', parents_enabled INTEGER NOT NULL DEFAULT 0);
CREATE TABLE IF NOT EXISTS enrollments(id INTEGER PRIMARY KEY, school_id INTEGER NOT NULL REFERENCES schools(id), student_id INTEGER NOT NULL REFERENCES users(id), class_id INTEGER NOT NULL REFERENCES classes(id), year_id INTEGER NOT NULL REFERENCES years(id), status TEXT NOT NULL DEFAULT 'inscrit', UNIQUE(student_id,class_id,year_id));
CREATE TABLE IF NOT EXISTS parent_links(id INTEGER PRIMARY KEY, school_id INTEGER NOT NULL REFERENCES schools(id), student_id INTEGER NOT NULL REFERENCES users(id), parent_id INTEGER NOT NULL REFERENCES users(id), relation TEXT DEFAULT 'parent', active INTEGER NOT NULL DEFAULT 1, UNIQUE(student_id,parent_id));
CREATE TABLE IF NOT EXISTS courses(id INTEGER PRIMARY KEY, school_id INTEGER NOT NULL REFERENCES schools(id), class_id INTEGER NOT NULL REFERENCES classes(id), name TEXT NOT NULL, code TEXT NOT NULL, planned_minutes INTEGER NOT NULL DEFAULT 1200, teacher_id INTEGER REFERENCES users(id), coefficient REAL NOT NULL DEFAULT 1);
CREATE TABLE IF NOT EXISTS teacher_history(id INTEGER PRIMARY KEY, course_id INTEGER NOT NULL REFERENCES courses(id), teacher_id INTEGER NOT NULL REFERENCES users(id), date_start TEXT NOT NULL, date_end TEXT);
CREATE TABLE IF NOT EXISTS lessons(id INTEGER PRIMARY KEY, school_id INTEGER NOT NULL REFERENCES schools(id), course_id INTEGER NOT NULL REFERENCES courses(id), teacher_id INTEGER NOT NULL REFERENCES users(id), starts_at TEXT NOT NULL, ends_at TEXT NOT NULL, room TEXT, status TEXT NOT NULL DEFAULT 'programmee', actual_minutes INTEGER NOT NULL DEFAULT 0, comment TEXT, validated INTEGER NOT NULL DEFAULT 0);
CREATE TABLE IF NOT EXISTS attendance(id INTEGER PRIMARY KEY, lesson_id INTEGER NOT NULL REFERENCES lessons(id), student_id INTEGER NOT NULL REFERENCES users(id), status TEXT NOT NULL, UNIQUE(lesson_id,student_id));
CREATE TABLE IF NOT EXISTS assessments(id INTEGER PRIMARY KEY, school_id INTEGER NOT NULL REFERENCES schools(id), course_id INTEGER NOT NULL REFERENCES courses(id), title TEXT NOT NULL, kind TEXT NOT NULL DEFAULT 'Devoir', due_date TEXT, coefficient REAL NOT NULL DEFAULT 1, max_score REAL NOT NULL DEFAULT 20, description TEXT, published INTEGER NOT NULL DEFAULT 1, attachment_name TEXT, attachment_blob BLOB, attachment_mime TEXT);
CREATE TABLE IF NOT EXISTS grades(id INTEGER PRIMARY KEY, assessment_id INTEGER NOT NULL REFERENCES assessments(id), student_id INTEGER NOT NULL REFERENCES users(id), value REAL NOT NULL, updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP, comment TEXT DEFAULT '', UNIQUE(assessment_id,student_id));
CREATE TABLE IF NOT EXISTS grade_history(id INTEGER PRIMARY KEY, grade_id INTEGER NOT NULL REFERENCES grades(id), old_value REAL, new_value REAL NOT NULL, actor_id INTEGER REFERENCES users(id), changed_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS periods(id INTEGER PRIMARY KEY, school_id INTEGER NOT NULL REFERENCES schools(id), class_id INTEGER NOT NULL REFERENCES classes(id), name TEXT NOT NULL, locked INTEGER NOT NULL DEFAULT 0, published INTEGER NOT NULL DEFAULT 0);
CREATE TABLE IF NOT EXISTS unlocks(id INTEGER PRIMARY KEY, period_id INTEGER NOT NULL REFERENCES periods(id), teacher_id INTEGER NOT NULL REFERENCES users(id), expires_at TEXT NOT NULL, reason TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS notifications(id INTEGER PRIMARY KEY, school_id INTEGER NOT NULL REFERENCES schools(id), user_id INTEGER NOT NULL REFERENCES users(id), message TEXT NOT NULL, href TEXT DEFAULT '', read INTEGER NOT NULL DEFAULT 0, created_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS push_subscriptions(id INTEGER PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE, endpoint TEXT NOT NULL UNIQUE, subscription_json TEXT NOT NULL, created_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS audit(id INTEGER PRIMARY KEY, school_id INTEGER, user_id INTEGER, action TEXT NOT NULL, details TEXT, created_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS submissions(id INTEGER PRIMARY KEY, assessment_id INTEGER NOT NULL REFERENCES assessments(id), student_id INTEGER NOT NULL REFERENCES users(id), filename TEXT NOT NULL, filepath TEXT NOT NULL, created_at TEXT DEFAULT CURRENT_TIMESTAMP, UNIQUE(assessment_id,student_id));
CREATE TABLE IF NOT EXISTS invites(id INTEGER PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id), token_hash TEXT NOT NULL, expires_at TEXT NOT NULL, used INTEGER NOT NULL DEFAULT 0);
CREATE TABLE IF NOT EXISTS mfa_settings(user_id INTEGER PRIMARY KEY REFERENCES users(id) ON DELETE CASCADE, totp_secret TEXT, totp_pending_secret TEXT, totp_last_step INTEGER NOT NULL DEFAULT -1, totp_attempts INTEGER NOT NULL DEFAULT 0, totp_locked_until INTEGER NOT NULL DEFAULT 0, email_code_hash TEXT, email_expires INTEGER, email_attempts INTEGER NOT NULL DEFAULT 0, email_last_sent INTEGER NOT NULL DEFAULT 0);
CREATE TABLE IF NOT EXISTS mfa_recovery_codes(user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE, code_hash TEXT NOT NULL, PRIMARY KEY(user_id,code_hash));
CREATE TABLE IF NOT EXISTS mfa_trusted_devices(id INTEGER PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE, token_hash TEXT NOT NULL UNIQUE, expires_at INTEGER NOT NULL, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS email_delivery_log(id INTEGER PRIMARY KEY, school_id INTEGER REFERENCES schools(id), recipient_user_id INTEGER REFERENCES users(id), sent_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS password_reset_limits(key_hash TEXT PRIMARY KEY, attempted_at INTEGER NOT NULL);
CREATE TABLE IF NOT EXISTS assessment_files(id INTEGER PRIMARY KEY, assessment_id INTEGER NOT NULL REFERENCES assessments(id) ON DELETE CASCADE, filename TEXT NOT NULL, mime TEXT NOT NULL, content BLOB NOT NULL);
CREATE TABLE IF NOT EXISTS bulletin_publications(id INTEGER PRIMARY KEY, school_id INTEGER NOT NULL REFERENCES schools(id), class_id INTEGER NOT NULL REFERENCES classes(id), student_id INTEGER NOT NULL REFERENCES users(id), published_at TEXT DEFAULT CURRENT_TIMESTAMP, UNIQUE(class_id,student_id));
CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY,value TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS matricule_rules(school_id INTEGER NOT NULL REFERENCES schools(id),role TEXT NOT NULL,prefix TEXT NOT NULL, next_number INTEGER NOT NULL DEFAULT 1, padding INTEGER NOT NULL DEFAULT 4,PRIMARY KEY(school_id,role));
"""

def connection():
    c=sqlite3.connect(DB,timeout=30)
    c.row_factory=sqlite3.Row
    c.execute('PRAGMA foreign_keys=ON')
    c.execute('PRAGMA busy_timeout=30000')
    return c

@contextmanager
def db():
    c=connection()
    try:
        yield c
        c.commit()
    except:
        c.rollback()
        raise
    finally:
        c.close()

def one(sql, args=()):
    with db() as c: return c.execute(sql,args).fetchone()

def allrows(sql,args=()):
    with db() as c: return c.execute(sql,args).fetchall()

def execute(sql,args=()):
    with db() as c:
        cur=c.execute(sql,args)
        return cur.lastrowid

Path(DB).parent.mkdir(parents=True,exist_ok=True)
with db() as c:
    c.executescript(SCHEMA)
    existing={r['name'] for r in c.execute('PRAGMA table_info(schools)').fetchall()}
    for name,definition in [('subscription_end','TEXT'),('logo','BLOB'),('logo_mime','TEXT'),('campus_enabled','INTEGER NOT NULL DEFAULT 0')]:
        if name not in existing: c.execute(f'ALTER TABLE schools ADD COLUMN {name} {definition}')
    columns={r['name'] for r in c.execute('PRAGMA table_info(users)').fetchall()}
    for name,definition in [('first_name','TEXT'),('last_name','TEXT'),('phone','TEXT'),('must_change_password','INTEGER NOT NULL DEFAULT 0'),('matricule','TEXT'),('first_login_at','TEXT')]:
        if name not in columns:c.execute(f'ALTER TABLE users ADD COLUMN {name} {definition}')
    for table, additions in {'courses':[('coefficient','REAL NOT NULL DEFAULT 1')],
        'assessments':[('attachment_name','TEXT'),('attachment_blob','BLOB'),('attachment_mime','TEXT')],
        'grades':[('comment', "TEXT DEFAULT ''")]}.items():
        existing_columns={x['name'] for x in c.execute(f'PRAGMA table_info({table})')}
        for column,definition in additions:
            if column not in existing_columns:c.execute(f'ALTER TABLE {table} ADD COLUMN {column} {definition}')
    c.execute('CREATE UNIQUE INDEX IF NOT EXISTS idx_users_school_email ON users(school_id,email) WHERE school_id IS NOT NULL')
    c.execute('CREATE UNIQUE INDEX IF NOT EXISTS idx_users_platform_email ON users(email) WHERE school_id IS NULL')
    c.execute('CREATE UNIQUE INDEX IF NOT EXISTS idx_users_school_matricule ON users(school_id,matricule) WHERE matricule IS NOT NULL')
    c.execute("INSERT OR IGNORE INTO settings(key,value) VALUES ('support_name','LMurbs')")
    c.execute("INSERT OR IGNORE INTO settings(key,value) VALUES ('support_phone','+2250711124359')")
    c.execute("INSERT OR IGNORE INTO settings(key,value) VALUES ('powered_by','Propulsé par LMurbs')")
MATRICULE_PROFILES = {'etudiant': 'ETU', 'enseignant': 'ENS', 'administratif': 'ADM', 'admin': 'DIR', 'parent': 'PAR'}

def ensure_matricule_rules(c, sid):
    for role, prefix in MATRICULE_PROFILES.items():
        c.execute('INSERT OR IGNORE INTO matricule_rules(school_id,role,prefix,next_number,padding) VALUES (?,?,?,1,4)', (sid,role,prefix))

def validate_matricule(value):
    value=(value or '').strip().upper()
    if not re.fullmatch(r'[A-Z0-9][A-Z0-9_/-]{1,39}',value):
        raise ValueError('Le matricule doit contenir 2 à 40 caractères (lettres, chiffres, tiret, / ou _).')
    return value

def next_matricule(c, sid, role, advance=True):
    ensure_matricule_rules(c,sid)
    rule=c.execute('SELECT prefix,next_number,padding FROM matricule_rules WHERE school_id=? AND role=?',(sid,role)).fetchone()
    if not rule:raise ValueError('Profil non configuré')
    number=rule['next_number']
    while True:
        candidate=f"{rule['prefix']}{number:0{rule['padding']}d}"
        if not c.execute('SELECT 1 FROM users WHERE school_id=? AND matricule=?',(sid,candidate)).fetchone():
            break
        number+=1
    if advance:
        c.execute('UPDATE matricule_rules SET next_number=? WHERE school_id=? AND role=?',(number+1,sid,role))
    return candidate

def requested_matricule(c,sid,role,raw):
    raw=(raw or '').strip()
    if raw:
        candidate=validate_matricule(raw)
        if c.execute('SELECT 1 FROM users WHERE school_id=? AND matricule=?',(sid,candidate)).fetchone():
            raise ValueError('Ce matricule existe déjà dans cet établissement.')
        if candidate==next_matricule(c,sid,role,advance=False):
            next_matricule(c,sid,role,advance=True)
        return candidate
    return next_matricule(c,sid,role)

# Migration non destructive : attribuer les numéros manquants aux comptes existants.
with db() as c:
    c.execute('BEGIN IMMEDIATE')
    for school in c.execute('SELECT id FROM schools').fetchall():
        sid=school['id']
        ensure_matricule_rules(c,sid)
        for user in c.execute("SELECT id,role FROM users WHERE school_id=? AND (matricule IS NULL OR matricule='') ORDER BY id",(sid,)).fetchall():
            if user['role'] in MATRICULE_PROFILES:
                c.execute('UPDATE users SET matricule=? WHERE id=?',(next_matricule(c,sid,user['role']),user['id']))

def setting(key,default=''):
    row=one('SELECT value FROM settings WHERE key=?',(key,))
    return row['value'] if row else default

def subscription_blocked(user=None):
    user=user or current_user()
    if not user or user['role']=='superadmin' or not user['school_id']:return False
    school=one('SELECT active,subscription_end FROM schools WHERE id=?',(user['school_id'],))
    if not school:return True
    return not school['active'] or bool(school['subscription_end'] and school['subscription_end']<=datetime.date.today().isoformat())

def read_logo(file):
    if not file or not file.filename:return None,None
    # Accepter les logos sources haute résolution (ex. PNG EduPilot de 2,1 Mo).
    raw=file.read(5*1024*1024+1)
    if not raw:return None,None
    if len(raw)>5*1024*1024:abort(400,'Logo supérieur à 5 Mo. Utilisez une image PNG, JPG ou WebP plus légère.')
    if raw.startswith(b'\x89PNG\r\n\x1a\n'):mime='image/png'
    elif raw[:3]==b'\xff\xd8\xff':mime='image/jpeg'
    elif raw[:4]==b'RIFF' and raw[8:12]==b'WEBP':mime='image/webp'
    else:abort(400,'Format de logo non reconnu.')
    return raw,mime

@app.get('/platform-logo')
def platform_logo():
    row=one('SELECT logo,logo_mime FROM platform_branding WHERE id=1')
    if not row or not row['logo']:abort(404)
    resp=Response(row['logo'],mimetype=row['logo_mime'])
    resp.headers['X-Content-Type-Options']='nosniff'
    resp.headers['Cache-Control']='public, max-age=60, must-revalidate'
    return resp

@app.get('/school-logo/<int:sid>')
def school_logo(sid):
    row=one('SELECT logo,logo_mime FROM schools WHERE id=?',(sid,))
    if not row or not row['logo']:abort(404)
    resp=Response(row['logo'],mimetype=row['logo_mime'])
    resp.headers['X-Content-Type-Options']='nosniff'
    resp.headers['Cache-Control']='public, max-age=300'
    return resp


@app.before_request
def safety():
    if request.method=='POST':
        token=request.form.get('csrf','') or request.headers.get('X-CSRF-Token','')
        if token!=session.get('csrf') or not token: abort(400,'Jeton CSRF invalide')
    if 'csrf' not in session: session['csrf']=secrets.token_urlsafe(24)
    u=current_user()
    if u and subscription_blocked(u) and request.endpoint not in ('restricted','logout','login','health','school_logo','static'):
        return redirect(url_for('restricted'))
    if u and u['must_change_password'] and request.endpoint not in ('change_password','logout','restricted','health','school_logo','static'):
        return redirect(url_for('change_password'))

@app.context_processor
def globals_template():
    return dict(me=current_user(),csrf=session.get('csrf',''),appname='EduPilot',support_name=setting('support_name','LMurbs'),support_phone=setting('support_phone','+2250711124359'),powered_by=setting('powered_by','Propulsé par LMurbs'),school_brand=one('SELECT id,name,logo_mime,campus_enabled FROM schools WHERE id=?',(school_id(),)) if school_id() else None, platform_logo_enabled=bool(one('SELECT id FROM platform_branding WHERE id=1')))

def current_user():
    uid=session.get('uid')
    return one('SELECT * FROM users WHERE id=? AND active=1',(uid,)) if uid else None

def login_required(fn):
    @functools.wraps(fn)
    def wrap(*args,**kwargs):
        if not current_user(): return redirect(url_for('login'))
        return fn(*args,**kwargs)
    return wrap

def roles(*allowed):
    def decorator(fn):
        @functools.wraps(fn)
        @login_required
        def wrap(*args,**kwargs):
            if current_user()['role'] not in allowed: abort(403)
            return fn(*args,**kwargs)
        return wrap
    return decorator

def school_id():
    u=current_user()
    return u['school_id'] if u else None

def scoped(sql,args=()):
    return allrows(sql,(school_id(),)+tuple(args))

def has_school(table,obj_id):
    return one(f'SELECT id FROM {table} WHERE school_id=? AND id=?',(school_id(),obj_id))

def audit(action,details=''):
    u=current_user()
    execute('INSERT INTO audit(school_id,user_id,action,details) VALUES (?,?,?,?)',(u['school_id'],u['id'],action,details))

def send_brevo_email(address,subject,message,school_id=None,recipient_user_id=None):
    """Email transactionnel Brevo. Ne journalise ni identifiants ni liens secrets."""
    api_key=os.getenv('BREVO_API_KEY','').strip()
    sender=os.getenv('BREVO_SENDER_EMAIL','').strip()
    if not api_key or not sender:
        app.logger.warning('Brevo non configuré : e-mail non envoyé')
        return False
    try:
        response=requests.post('https://api.brevo.com/v3/smtp/email',
            json={'sender':{'email':sender,'name':os.getenv('BREVO_SENDER_NAME','EduPilot')},
                  'to':[{'email':address}],'subject':subject,'textContent':message},
            headers={'api-key':api_key,'Content-Type':'application/json'},timeout=8)
        response.raise_for_status()
        execute('INSERT INTO email_delivery_log(school_id,recipient_user_id) VALUES (?,?)',(school_id,recipient_user_id))
        return True
    except requests.RequestException as exc:
        app.logger.warning('Échec e-mail Brevo : %s',type(exc).__name__)
        return False

def send_account_link(uid,purpose='activation'):
    """Émet un lien à usage unique, valable 48 h pour activation, 30 min pour réinitialisation."""
    user=one('SELECT u.id,u.name,u.email,u.active,u.role,u.school_id,s.name school_name FROM users u LEFT JOIN schools s ON s.id=u.school_id WHERE u.id=?',(uid,))
    if not user or not user['active']:return False
    token=secrets.token_urlsafe(32)
    token_hash=hashlib.sha256(token.encode()).hexdigest()
    expiration=datetime.datetime.now(datetime.timezone.utc)+datetime.timedelta(minutes=30 if purpose=='reset' else 48*60)
    with db() as c:
        c.execute('DELETE FROM invites WHERE user_id=? AND used=0',(uid,))
        c.execute('INSERT INTO invites(user_id,token_hash,expires_at) VALUES (?,?,?)',(uid,token_hash,expiration.isoformat()))
    base=os.getenv('PUBLIC_BASE_URL','https://edupilot.ci').rstrip('/')
    url=base+url_for('activate',token=token)
    if purpose=='reset':
        subject='EduPilot — Réinitialisation de votre mot de passe'
        content=f"Bonjour {user['name']},\n\nPour définir un nouveau mot de passe EduPilot, ouvrez ce lien personnel :\n{url}\n\nCe lien expire dans 30 minutes et ne peut servir qu'une fois. Si vous n'avez rien demandé, ignorez ce message.\n\nEduPilot"
    else:
        subject='Bienvenue sur EduPilot — Activez votre compte'
        content=f"Bonjour {user['name']},\n\nVotre compte EduPilot ({user['role']}) pour {user['school_name'] or 'la plateforme'} a été créé.\n\nPour choisir votre mot de passe et activer votre accès :\n{url}\n\nCe lien personnel expire dans 48 heures. Ne le transférez à personne.\n\nPour utiliser EduPilot sur votre téléphone ou tablette, consultez le guide d'installation :\n{base}/installer\n\nAprès installation, connectez-vous avec votre compte EduPilot.\n\nEduPilot"
    sent=send_brevo_email(user['email'],subject,content,user['school_id'],uid)
    # Ne pas enregistrer les URL contenant un jeton dans les notifications internes.
    return sent

def notify(ids,message,href=''):
    # Persist first: an external push or email error cannot erase an internal alert.
    sid=school_id()
    candidate=set(ids)
    recipients=[r['id'] for r in allrows('SELECT id FROM users WHERE school_id=? AND active=1',(sid,)) if r['id'] in candidate]
    for uid in recipients:
        execute('INSERT INTO notifications(school_id,user_id,message,href) VALUES (?,?,?,?)',(sid,uid,message,href))
    if recipients and all(os.getenv(k) for k in ('VAPID_PRIVATE_KEY','VAPID_PUBLIC_KEY','VAPID_SUBJECT')):
        try:
            from pywebpush import webpush, WebPushException
            subscriptions=allrows('SELECT id,user_id,subscription_json FROM push_subscriptions WHERE user_id IN ('+','.join('?' for _ in recipients)+')',tuple(recipients))
            for subscription in subscriptions:
                try:
                    webpush(subscription_info=json.loads(subscription['subscription_json']),
                        data=json.dumps({'title':'EduPilot','body':message[:350],'url':href if href.startswith('/') and not href.startswith('//') else '/'}),
                        vapid_private_key=os.getenv('VAPID_PRIVATE_KEY'),
                        vapid_claims={'sub':os.getenv('VAPID_SUBJECT')},timeout=3)
                except WebPushException as exc:
                    if getattr(getattr(exc,'response',None),'status_code',None) in (404,410):
                        execute('DELETE FROM push_subscriptions WHERE id=?',(subscription['id'],))
                    else:app.logger.warning('Échec push: %s',type(exc).__name__)
                except Exception as exc:app.logger.warning('Échec push: %s',type(exc).__name__)
        except ImportError:app.logger.warning('Module pywebpush indisponible')
    for uid in recipients:
        user=one('SELECT email FROM users WHERE id=? AND school_id=?',(uid,sid))
        if user:send_brevo_email(user['email'],'EduPilot — Nouvelle notification',message,sid,uid)


def classmates(course_id):
    return [r['student_id'] for r in scoped('SELECT e.student_id FROM enrollments e JOIN courses c ON c.class_id=e.class_id WHERE e.school_id=? AND c.id=?',(course_id,))]

def event_recipients(course_ids=(),teacher_ids=(),include_management=True,include_parents=True):
    """Recipients scoped to the current establishment and relevant classes."""
    sid=school_id()
    recipients=set()
    courses=[int(cid) for cid in course_ids if cid]
    teachers=[int(uid) for uid in teacher_ids if uid]
    if courses:
        clause=','.join('?' for _ in courses)
        students=allrows(f"""SELECT DISTINCT e.student_id FROM enrollments e
            JOIN courses c ON c.class_id=e.class_id AND c.school_id=e.school_id
            JOIN users u ON u.id=e.student_id
            WHERE e.school_id=? AND c.id IN ({clause}) AND u.active=1""",(sid,*courses))
        recipients.update(r['student_id'] for r in students)
        if include_parents:
            parents=allrows(f"""SELECT DISTINCT p.parent_id FROM parent_links p
                JOIN enrollments e ON e.student_id=p.student_id AND e.school_id=p.school_id
                JOIN courses c ON c.class_id=e.class_id AND c.school_id=e.school_id
                JOIN classes cl ON cl.id=c.class_id
                JOIN users u ON u.id=p.parent_id
                WHERE p.school_id=? AND c.id IN ({clause})
                AND p.active=1 AND cl.parents_enabled=1 AND u.active=1""",(sid,*courses))
            recipients.update(r['parent_id'] for r in parents)
    if teachers:
        clause=','.join('?' for _ in teachers)
        selected=allrows(f"""SELECT id FROM users WHERE school_id=? AND active=1
            AND role IN ('enseignant','admin') AND id IN ({clause})""",(sid,*teachers))
        recipients.update(r['id'] for r in selected)
    if include_management:
        managers=allrows("""SELECT id FROM users WHERE school_id=? AND active=1
                            AND role IN ('admin','administratif')""",(sid,))
        recipients.update(r['id'] for r in managers)
    return sorted(recipients)

def notify_event(course_ids,message,href,teacher_ids=(),include_parents=True):
    recipients=event_recipients(course_ids,teacher_ids,include_parents=include_parents)
    # No duplicate notification to the user who performed the operation.
    actor=current_user()
    if actor:
        recipients=[uid for uid in recipients if uid!=actor['id']]
    if recipients:
        notify(recipients,message,href)


@app.template_filter('statut_seance')
def statut_seance(value):
    return {'programmee':'Programmée','realisee':'Réalisée','partielle':'Partiellement réalisée',
            'non_effectuee':'Non effectuée','annulee':'Annulée','reportee':'Reportée'}.get(value,value.replace('_',' ').capitalize() if value else '—')

def slugify(s):
    s=unicodedata.normalize('NFKD',s).encode('ascii','ignore').decode().lower()
    return re.sub(r'^_+|_+$','',re.sub(r'[^a-z0-9]+','_',s))

@app.get('/health')
def health():return {'status':'ok','application':'EduPilot'}

@app.get('/<slug>')
@app.get('/<slug>/login')
def tenant_login(slug):
    school=one('SELECT id,name,slug,logo_mime FROM schools WHERE slug=?',(slug,))
    if not school:abort(404)
    return render_template('login.html',login_school=school,login_action=url_for('tenant_signin',slug=slug))

@app.post('/<slug>/login')
def tenant_signin(slug):
    school=one('SELECT id,name,slug,logo_mime FROM schools WHERE slug=?',(slug,))
    if not school:abort(404)
    return handle_login(school)

def complete_login(u, trusted_cookie=None):
    # A successful authentication, not opening an invitation link, validates it.
    execute("UPDATE users SET first_login_at=CURRENT_TIMESTAMP WHERE id=? AND first_login_at IS NULL",(u['id'],))
    session.clear()
    session['uid']=u['id']
    session['csrf']=secrets.token_urlsafe(24)
    if u['must_change_password']:
        return redirect(url_for('change_password'))
    response=make_response(redirect(url_for('dashboard')))
    if trusted_cookie:
        response.set_cookie('edupilot_trusted',trusted_cookie,max_age=(30 if u['role'] in ('admin','administratif') else 90)*86400,secure=True,httponly=True,samesite='Lax',path='/')
    return response


def mfa_role(u):
    if u['role']=='superadmin':return 'totp'
    if u['role'] in ('admin','administratif','etudiant','enseignant','parent'):return 'email'
    return None


def generate_recovery_codes(uid):
    codes=['-'.join((secrets.token_hex(4),secrets.token_hex(4))) for _ in range(8)]
    with db() as c:
        c.execute('DELETE FROM mfa_recovery_codes WHERE user_id=?',(uid,))
        c.executemany('INSERT INTO mfa_recovery_codes(user_id,code_hash) VALUES (?,?)',
                      [(uid,hashlib.sha256(code.encode()).hexdigest()) for code in codes])
    return codes


def use_recovery_code(uid,code):
    code=code.strip().lower()
    if not re.fullmatch(r'[0-9a-f]{8}-[0-9a-f]{8}',code):return False
    with db() as c:
        removed=c.execute('DELETE FROM mfa_recovery_codes WHERE user_id=? AND code_hash=?',
                          (uid,hashlib.sha256(code.encode()).hexdigest()))
        return removed.rowcount==1


def totp_qr_data(uri):
    import qrcode
    from io import BytesIO
    image=qrcode.make(uri,box_size=6,border=3)
    output=BytesIO()
    image.save(output,format='PNG')
    return 'data:image/png;base64,'+base64.b64encode(output.getvalue()).decode()


def mfa_valid_totp(secret,code,previous=-1):
    try:
        key=base64.b32decode(secret+'='*((-len(secret))%8))
        if not (len(code)==6 and code.isascii() and code.isdigit()):return None
        current=int(time.time())//30
        for step in (current-1,current,current+1):
            if step<=previous:continue
            digest=hmac.new(key,struct.pack('>Q',step),hashlib.sha1).digest()
            p=digest[-1]&15
            value=(struct.unpack('>I',digest[p:p+4])[0]&0x7fffffff)%1000000
            if hmac.compare_digest(f'{value:06d}',code):return step
    except (ValueError,TypeError):pass
    return None


def mfa_trusted_user(u):
    token=request.cookies.get('edupilot_trusted','')
    if not token or len(token)>180:return False
    digest=hashlib.sha256(token.encode()).hexdigest()
    return bool(one('SELECT id FROM mfa_trusted_devices WHERE user_id=? AND token_hash=? AND expires_at>?',(u['id'],digest,int(time.time()))))


def start_second_factor(u):
    mode=mfa_role(u)
    if not mode or (mode=='email' and mfa_trusted_user(u)):
        return complete_login(u)
    session.clear()
    session['csrf']=secrets.token_urlsafe(24)
    session['mfa_pending_uid']=u['id']
    session['mfa_pending_at']=int(time.time())
    if mode=='email':
        send_mfa_email(u)
    return redirect(url_for('mfa_verify'))


def send_mfa_email(u):
    now=int(time.time())
    with db() as c:
        c.execute('INSERT OR IGNORE INTO mfa_settings(user_id) VALUES (?)',(u['id'],))
        row=c.execute('SELECT email_last_sent FROM mfa_settings WHERE user_id=?',(u['id'],)).fetchone()
        if now-row['email_last_sent']<60:return False
        code=f'{secrets.randbelow(1000000):06d}'
        digest=hmac.new(app.secret_key.encode(),f"{u['id']}:{code}".encode(),hashlib.sha256).hexdigest()
        c.execute('UPDATE mfa_settings SET email_code_hash=?,email_expires=?,email_attempts=0,email_last_sent=? WHERE user_id=?',(digest,now+600,now,u['id']))
    return send_brevo_email(u['email'],'EduPilot — Code de connexion',f'Bonjour,\n\nVotre code de vérification EduPilot est : {code}\n\nCe code est valable pendant 10 minutes.\n\nSi vous n’êtes pas à l’origine de cette demande, ignorez cet e-mail.\n\nL’équipe EduPilot',u['school_id'],u['id'])


@app.route('/mfa',methods=['GET','POST'])
def mfa_verify():
    uid=session.get('mfa_pending_uid')
    stamp=session.get('mfa_pending_at',0)
    u=one('SELECT * FROM users WHERE id=? AND active=1',(uid,)) if uid and int(time.time())-stamp<600 else None
    if not u:return redirect(url_for('login'))
    mode=mfa_role(u)
    if mode not in ('totp','email'):abort(403)
    config=one('SELECT * FROM mfa_settings WHERE user_id=?',(u['id'],))
    if mode=='totp' and (not config or not config['totp_secret']):
        if not config or not config['totp_pending_secret']:
            candidate=base64.b32encode(secrets.token_bytes(20)).decode().rstrip('=')
            with db() as c:
                c.execute('INSERT OR IGNORE INTO mfa_settings(user_id) VALUES (?)',(u['id'],))
                c.execute('UPDATE mfa_settings SET totp_pending_secret=? WHERE user_id=? AND totp_secret IS NULL',(candidate,u['id']))
            config=one('SELECT * FROM mfa_settings WHERE user_id=?',(u['id'],))
        secret=config['totp_pending_secret']
        uri='otpauth://totp/'+quote('EduPilot:'+u['email'])+'?secret='+secret+'&issuer=EduPilot&digits=6&period=30'
    else:
        secret=None
        uri=None
    if request.method=='POST':
        if mode=='totp':
            if config and config['totp_locked_until']>int(time.time()):
                flash('Trop de tentatives. Réessayez plus tard.','error')
                return redirect(url_for('mfa_verify'))
            with db() as c:
                c.execute('UPDATE mfa_settings SET totp_attempts=totp_attempts+1 WHERE user_id=?',(u['id'],))
                c.execute('UPDATE mfa_settings SET totp_locked_until=?,totp_attempts=0 WHERE user_id=? AND totp_attempts>=5',(int(time.time())+300,u['id']))
        if mode=='totp' and request.form.get('action')=='email_send' and config and config['totp_secret']:
            if send_mfa_email(u):
                session['mfa_email_fallback']=True
                flash('Un code de secours a été envoyé à votre adresse e-mail.','success')
            else:
                flash('Envoi indisponible ou délai entre deux demandes non écoulé.','error')
            return redirect(url_for('mfa_verify'))
        if mode=='email' and request.form.get('action')=='resend':
            send_mfa_email(u)
            flash('Si l’envoi est disponible, un nouveau code a été envoyé.','success')
            return redirect(url_for('mfa_verify'))
        code=request.form.get('code','').strip()
        if mode=='totp' and config and config['totp_secret'] and session.get('mfa_email_fallback'):
            if config['email_expires'] and config['email_expires']>=int(time.time()) and config['email_attempts']<5 and len(code)==6:
                digest=hmac.new(app.secret_key.encode(),f"{u['id']}:{code}".encode(),hashlib.sha256).hexdigest()
                with db() as c:
                    c.execute('UPDATE mfa_settings SET email_attempts=email_attempts+1 WHERE user_id=?',(u['id'],))
                    updated=c.execute('UPDATE mfa_settings SET email_code_hash=NULL,email_expires=NULL WHERE user_id=? AND email_code_hash=? AND email_expires>=? AND email_attempts<=5',(u['id'],digest,int(time.time())))
                if updated.rowcount==1:
                    execute('INSERT INTO audit(school_id,user_id,action,details) VALUES (?,?,?,?)',(u['school_id'],u['id'],'mfa.superadmin.email_fallback','Second facteur vérifié par e-mail'))
                    return complete_login(u)
        if mode=='totp' and config and config['totp_secret'] and use_recovery_code(u['id'],code):
            return complete_login(u)
        if mode=='totp':
            step=mfa_valid_totp(secret or config['totp_secret'],code,config['totp_last_step'] if config else -1)
            if step is not None:
                with db() as c:
                    c.execute('INSERT OR IGNORE INTO mfa_settings(user_id) VALUES (?)',(u['id'],))
                    if secret:
                        c.execute('UPDATE mfa_settings SET totp_secret=?,totp_pending_secret=NULL,totp_last_step=?,totp_attempts=0,totp_locked_until=0 WHERE user_id=?',(secret,step,u['id']))
                    else:
                        changed=c.execute('UPDATE mfa_settings SET totp_last_step=?,totp_attempts=0,totp_locked_until=0 WHERE user_id=? AND totp_last_step<?',(step,u['id'],step))
                        if changed.rowcount!=1:abort(403)
                if secret:
                    codes=generate_recovery_codes(u['id'])
                    session.clear()
                    session['uid']=u['id']
                    session['csrf']=secrets.token_urlsafe(24)
                    execute("UPDATE users SET first_login_at=CURRENT_TIMESTAMP WHERE id=? AND first_login_at IS NULL",(u['id'],))
                    return render_template('mfa_recovery_codes.html',codes=codes)
                return complete_login(u)
        elif config and config['email_expires'] and config['email_expires']>=int(time.time()) and config['email_attempts']<5 and len(code)==6:
            digest=hmac.new(app.secret_key.encode(),f"{u['id']}:{code}".encode(),hashlib.sha256).hexdigest()
            with db() as c:
                c.execute('UPDATE mfa_settings SET email_attempts=email_attempts+1 WHERE user_id=?',(u['id'],))
                if hmac.compare_digest(config['email_code_hash'] or '',digest):
                    c.execute('UPDATE mfa_settings SET email_code_hash=NULL,email_expires=NULL WHERE user_id=?',(u['id'],))
                    cookie=None
                    if request.form.get('trust_device')=='1':
                        cookie=secrets.token_urlsafe(32)
                        days=30 if u['role'] in ('admin','administratif') else 90
                        c.execute('DELETE FROM mfa_trusted_devices WHERE user_id=? AND expires_at<?',(u['id'],int(time.time())))
                        c.execute('INSERT INTO mfa_trusted_devices(user_id,token_hash,expires_at) VALUES (?,?,?)',(u['id'],hashlib.sha256(cookie.encode()).hexdigest(),int(time.time())+days*86400))
                    verified_cookie=cookie
        if mode=='email' and 'verified_cookie' in locals():
            return complete_login(u,verified_cookie)
        flash('Code incorrect ou expiré.','error')
    return render_template('mfa.html',mode=mode,setup_secret=secret,setup_uri=uri,setup_qr=totp_qr_data(uri) if uri else None,email_fallback=bool(session.get('mfa_email_fallback')) if mode=='totp' else False,trust_days=30 if u['role'] in ('admin','administratif') else 90)


def handle_login(school=None):
    email=request.form.get('email','').strip().lower()
    password=request.form.get('password','')
    if school:
        candidates=allrows('SELECT * FROM users WHERE email=? AND school_id=? AND active=1',(email,school['id']))
    else:
        candidates=allrows('SELECT * FROM users WHERE email=? AND active=1',(email,))
    valid=[u for u in candidates if check_password_hash(u['password_hash'],password)]
    if school:
        valid=[u for u in valid if u['role']!='superadmin']
    if len(valid)==1:
        return start_second_factor(valid[0])
    if not school and len(valid)>1:
        ids=[u['id'] for u in valid if u['school_id'] is not None]
        schools=allrows('SELECT s.id,s.name,s.slug FROM schools s JOIN users u ON u.school_id=s.id WHERE u.id IN ('+','.join('?' for _ in ids)+') AND s.active=1 ORDER BY s.name',tuple(ids)) if ids else []
        if len(schools)>1:
            # Only offer schools for which the supplied password was verified.
            session['pending_school_choice']={'ids':ids,'at':int(time.time())}
            return render_template('login.html',login_school=None,login_action=url_for('login'),school_choices=schools)
    flash('Identifiants incorrects.','error')
    if school:return redirect(url_for('tenant_login',slug=school['slug']))
    return redirect(url_for('login'))


@app.post('/login/choose-school')
def choose_login_school():
    pending=session.pop('pending_school_choice',None)
    try:
        school_id=int(request.form.get('school_id',''))
    except (ValueError,TypeError):
        abort(400)
    if not pending or int(time.time())-pending.get('at',0)>120:
        flash('Sélection expirée. Veuillez vous reconnecter.','error')
        return redirect(url_for('login'))
    ids=pending.get('ids',[])
    if not ids:
        abort(403)
    u=one('SELECT * FROM users WHERE school_id=? AND id IN ('+','.join('?' for _ in ids)+') AND active=1',(school_id,*ids))
    if not u or u['role']=='superadmin' or not one('SELECT id FROM schools WHERE id=? AND active=1',(school_id,)):
        abort(403)
    return start_second_factor(u)


@app.route('/login',methods=['GET','POST'])
def login():
    if request.method=='POST':return handle_login()
    session.pop('pending_school_choice',None)
    return render_template('login.html',login_school=None,login_action=url_for('login'))


@app.route('/change-password',methods=['GET','POST'])
@login_required
def change_password():
    u=current_user()
    if request.method=='POST':
        old=request.form.get('old_password','')
        new=request.form.get('new_password','')
        confirm=request.form.get('confirm_password','')
        if not check_password_hash(u['password_hash'],old):
            flash('Mot de passe temporaire incorrect.','error')
        elif len(new)<12 or not any(x.isalpha() for x in new) or not any(x.isdigit() for x in new):
            flash('Choisissez au moins 12 caractères, avec des lettres et des chiffres.','error')
        elif new!=confirm:
            flash('Les nouveaux mots de passe sont différents.','error')
        elif new==old:
            flash('Le nouveau mot de passe doit être différent.','error')
        else:
            execute('UPDATE users SET password_hash=?, must_change_password=0 WHERE id=?',(generate_password_hash(new),u['id']))
            session.clear()
            flash('Mot de passe modifié. Connectez-vous avec votre nouveau mot de passe.','success')
            return redirect(url_for('login'))
    return render_template('change_password.html')

@app.get('/restricted')
@login_required
def restricted():
    if not subscription_blocked():return redirect(url_for('dashboard'))
    return render_template('restricted.html')

@app.post('/logout')
def logout():
    # Revoke this device only, not all phones subscribed by the same user.
    uid=session.get('uid')
    endpoint=request.form.get('push_endpoint','')
    if uid and endpoint and len(endpoint)<=2048:
        execute('DELETE FROM push_subscriptions WHERE user_id=? AND endpoint=?',(uid,endpoint))
    session.clear()
    return redirect(url_for('login'))

@app.get('/')
@login_required
def dashboard():
    u=current_user()
    if u['role']=='superadmin':
        schools=allrows('SELECT s.*, (SELECT COUNT(*) FROM users u WHERE u.school_id=s.id) members FROM schools s')
        return render_template('dashboard.html',schools=schools,stats=None,classes=[],events=[],today=datetime.date.today().isoformat(),base_url=os.getenv('PUBLIC_BASE_URL',request.host_url.rstrip('/')).rstrip('/'))
    if u['role']=='etudiant':
        # Week selection is server-side and restricted to this student's own enrollments.
        ref=request.args.get('week','')
        try:
            today=datetime.date.fromisoformat(ref) if ref else datetime.date.today()
        except ValueError:
            today=datetime.date.today()
        # Dashboard = rolling seven-day preview, never show dates before today.
        # Full /planning calendar keeps complete historical navigation.
        today=max(today,datetime.date.today())
        monday=today
        sunday=today+datetime.timedelta(days=7)
        week_start=today.isoformat()
        week_end=sunday.isoformat()
        lessons=allrows("""SELECT l.id,l.starts_at,l.ends_at,l.room,l.status,c.name course_name,cl.name class_name,t.name teacher_name
                         FROM lessons l JOIN courses c ON c.id=l.course_id
                         JOIN classes cl ON cl.id=c.class_id LEFT JOIN users t ON t.id=l.teacher_id
                         WHERE l.school_id=? AND l.ends_at>? AND l.starts_at<?
                         AND EXISTS (SELECT 1 FROM enrollments e
                           WHERE e.school_id=l.school_id AND e.student_id=? AND e.class_id=c.class_id)
                         ORDER BY l.starts_at,c.name""",(u['school_id'],max(datetime.datetime.now().strftime('%Y-%m-%dT%H:%M'),week_start),week_end,u['id']))
        week_days=[{'date':monday+datetime.timedelta(days=n),'lessons':[]} for n in range(7)]
        for item in lessons:
            index=(datetime.date.fromisoformat(item['starts_at'][:10])-monday).days
            if 0<=index<7:week_days[index]['lessons'].append(item)
        return render_template('dashboard_student.html',today=datetime.date.today(),week_days=week_days,
                               monday=monday,prev_week=(monday-datetime.timedelta(days=7)).isoformat(),
                               next_week=(monday+datetime.timedelta(days=7)).isoformat(),
                               week_end=monday+datetime.timedelta(days=6),lesson_count=len(lessons))
    if u['role']=='parent':
        return render_template('dashboard.html',stats={'classes':0,'courses':0,'lessons':0,'total':0,'actual':0,'percent':0},classes=[],events=[],courses=[],schools=[])
    sid=school_id()
    classes=scoped('SELECT * FROM classes WHERE school_id=?')
    courses=scoped('SELECT c.*,cl.name class_name,u.name teacher_name,u.matricule teacher_matricule FROM courses c JOIN classes cl ON cl.id=c.class_id LEFT JOIN users u ON u.id=c.teacher_id WHERE c.school_id=?')
    if u['role']=='enseignant': courses=[x for x in courses if x['teacher_id']==u['id']]
    lessonlist=scoped('SELECT l.*,CAST(ROUND((julianday(l.ends_at)-julianday(l.starts_at))*1440) AS INTEGER) scheduled_minutes,c.name course_name,cl.name class_name,t.name teacher_name FROM lessons l JOIN courses c ON c.id=l.course_id JOIN classes cl ON cl.id=c.class_id LEFT JOIN users t ON t.id=l.teacher_id WHERE l.school_id=? ORDER BY l.starts_at DESC LIMIT 80')
    if u['role']=='enseignant': lessonlist=[x for x in lessonlist if any(y['id']==x['course_id'] for y in courses)]
    upcoming=scoped("""SELECT l.*,c.name course_name,cl.name class_name
        FROM lessons l JOIN courses c ON c.id=l.course_id
        JOIN classes cl ON cl.id=c.class_id
        WHERE l.school_id=? AND l.starts_at>=? AND l.status='programmee'
        ORDER BY l.starts_at ASC LIMIT 80""",(datetime.datetime.now().strftime('%Y-%m-%dT%H:%M'),))
    if u['role']=='enseignant':
        upcoming=[x for x in upcoming if x['teacher_id']==u['id']]
    upcoming=upcoming[:8]
    total=sum(x['planned_minutes'] for x in courses)
    actual=sum(x['actual_minutes'] for x in lessonlist if x['status'] in ('realisee','partielle'))
    stats={'classes':len(classes),'courses':len(courses),'lessons':len(lessonlist),'total':total//60,'actual':round(actual/60,1),'percent':round(actual/total*100,1) if total else 0}
    return render_template('dashboard.html',stats=stats,classes=classes,events=upcoming,schools=[],courses=courses)

@app.get('/schools')
@roles('superadmin')
def schools():
    records=allrows('SELECT id,name,slug,plan,active,subscription_end,logo_mime FROM schools ORDER BY id DESC')
    counts=allrows("SELECT school_id,COUNT(*) n FROM users WHERE role IN ('admin','administratif') GROUP BY school_id")
    account_counts={r['school_id']:r['n'] for r in counts}
    admins=allrows("SELECT school_id,first_name,last_name,name,email,phone,role,active,must_change_password FROM users WHERE role IN ('admin','administratif') ORDER BY id DESC")
    return render_template('schools.html',schools=records,admin_accounts=admins,account_counts=account_counts,today=datetime.date.today().isoformat(),base_url=os.getenv('PUBLIC_BASE_URL',request.host_url.rstrip('/')).rstrip('/'))

@app.get('/superadmin/users-report')
@roles('superadmin')
def superadmin_users_report():
    schools=allrows('SELECT id,name FROM schools ORDER BY name')
    years=allrows('SELECT id,school_id,label FROM years ORDER BY label DESC')
    classes=allrows('SELECT id,school_id,year_id,name FROM classes ORDER BY name')
    values={}
    allowed={'school_id':{r['id'] for r in schools},'year_id':{r['id'] for r in years},'class_id':{r['id'] for r in classes}}
    for key in ('school_id','year_id','class_id'):
        value=request.args.get(key,'').strip()
        if value:
            if not value.isdigit() or int(value) not in allowed[key]:abort(400,'Filtre invalide')
            values[key]=int(value)
    if 'year_id' in values and 'school_id' in values and not any(r['id']==values['year_id'] and r['school_id']==values['school_id'] for r in years):abort(400)
    if 'class_id' in values and not any(r['id']==values['class_id'] and ('school_id' not in values or r['school_id']==values['school_id']) and ('year_id' not in values or r['year_id']==values['year_id']) for r in classes):abort(400)
    where=["u.school_id IS NOT NULL"]
    args=[]
    if 'school_id' in values:where.append('u.school_id=?');args.append(values['school_id'])
    if 'year_id' in values:
        where.append("""(EXISTS (SELECT 1 FROM enrollments e WHERE e.student_id=u.id AND e.school_id=u.school_id AND e.year_id=?)
        OR EXISTS (SELECT 1 FROM courses c JOIN classes cl ON cl.id=c.class_id WHERE c.teacher_id=u.id AND c.school_id=u.school_id AND cl.year_id=?)
        OR EXISTS (SELECT 1 FROM parent_links p JOIN enrollments e ON e.student_id=p.student_id AND e.school_id=p.school_id WHERE p.parent_id=u.id AND p.school_id=u.school_id AND e.year_id=?))""")
        args.extend([values['year_id']]*3)
    if 'class_id' in values:
        where.append("""(EXISTS (SELECT 1 FROM enrollments e WHERE e.student_id=u.id AND e.school_id=u.school_id AND e.class_id=?)
        OR EXISTS (SELECT 1 FROM courses c WHERE c.teacher_id=u.id AND c.school_id=u.school_id AND c.class_id=?)
        OR EXISTS (SELECT 1 FROM parent_links p JOIN enrollments e ON e.student_id=p.student_id AND e.school_id=p.school_id WHERE p.parent_id=u.id AND p.school_id=u.school_id AND e.class_id=?))""")
        args.extend([values['class_id']]*3)
    users=allrows('SELECT u.id,u.name,u.email,u.role,u.active,u.school_id,s.name school_name FROM users u JOIN schools s ON s.id=u.school_id WHERE '+' AND '.join(where)+' ORDER BY s.name,u.role,u.name',tuple(args))
    counts={role:sum(1 for u in users if u['role']==role) for role in ('admin','administratif','enseignant','etudiant','parent')}
    mail_where=['1=1'];mail_args=[]
    if 'school_id' in values:mail_where.append('m.school_id=?');mail_args.append(values['school_id'])
    if 'year_id' in values or 'class_id' in values:
        ids=[u['id'] for u in users]
        if ids:mail_where.append('m.recipient_user_id IN ('+','.join('?' for _ in ids)+')');mail_args+=ids
        else:mail_where.append('1=0')
    mailed=one('SELECT COUNT(*) n FROM email_delivery_log m WHERE '+' AND '.join(mail_where),tuple(mail_args))['n']
    return render_template('superadmin_users_report.html',schools=schools,years=years,classes=classes,filters=values,users=users,counts=counts,mailed=mailed)


@app.route('/schools/new',methods=['GET','POST'])
@roles('superadmin')
def school_new():
    if request.method=='GET':return render_template('school_new.html')
    name=request.form.get('name','').strip()
    if not name:abort(400)
    end=request.form.get('subscription_end','').strip() or None
    if end:
        try:datetime.date.fromisoformat(end)
        except ValueError:abort(400)
    plan=request.form.get('plan','standard')
    if plan not in ('essentiel','standard','premium'):abort(400)
    slug=slugify(name)
    if not slug:abort(400)
    n=2;unique=slug
    while one('SELECT id FROM schools WHERE slug=?',(unique,)):
        unique=slug+'_'+str(n);n+=1
    logo,mime=read_logo(request.files.get('logo'))
    with db() as c:
        sid=c.execute('INSERT INTO schools(name,slug,plan,subscription_end,logo,logo_mime,campus_enabled) VALUES (?,?,?,?,?,?,?)',(name,unique,plan,end,logo,mime,int(request.form.get('campus_enabled')=='1'))).lastrowid
        c.execute('INSERT INTO years(school_id,label) VALUES (?,?)',(sid,'2026-2027'))
        ensure_matricule_rules(c,sid)
    flash('Établissement créé. Créez maintenant ses comptes administrateurs.','success')
    return redirect(url_for('school_accounts',sid=sid))

@app.route('/schools/<int:sid>/accounts',methods=['GET','POST'])
@roles('superadmin')
def school_accounts(sid):
    school=one('SELECT id,name,slug,logo_mime,plan,active,subscription_end FROM schools WHERE id=?',(sid,))
    if not school:abort(404)
    if request.method=='POST':
        first=request.form.get('first_name','').strip()
        last=request.form.get('last_name','').strip()
        email=request.form.get('email','').strip().lower()
        phone=request.form.get('phone','').strip()
        role=request.form.get('role','')
        password=secrets.token_urlsafe(32)
        if not first or not last or not phone or not email or not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+',email):abort(400,'Informations obligatoires manquantes')
        if len(first)>90 or len(last)>90 or len(phone)>40 or len(email)>254:abort(400)
        if role not in ('admin','administratif'):abort(400)
        if one('SELECT id FROM users WHERE school_id=? AND email=?',(sid,email)):
            flash('Cette adresse e-mail est déjà utilisée.','error')
        else:
            try:
                with db() as c:
                    c.execute('BEGIN IMMEDIATE')
                    matricule=requested_matricule(c,sid,role,request.form.get('matricule'))
                    uid=c.execute('INSERT INTO users(school_id,name,first_name,last_name,email,phone,password_hash,role,matricule,must_change_password) VALUES (?,?,?,?,?,?,?,?,?,1)',(sid,first+' '+last,first,last,email,phone,generate_password_hash(password),role,matricule)).lastrowid
            except (ValueError,sqlite3.IntegrityError) as exc:
                flash(str(exc),'error')
                return redirect(url_for('school_accounts',sid=sid))
            sent=create_invite(uid)
            flash('Compte créé. '+('Invitation envoyée par e-mail.' if sent else 'E-mail non envoyé : vérifiez Brevo et renvoyez l’invitation.'),'success' if sent else 'error')
            return redirect(url_for('school_accounts',sid=sid))
    accounts=allrows("SELECT id,first_name,last_name,name,email,phone,role,matricule,active,must_change_password,first_login_at FROM users WHERE school_id=? AND role IN ('admin','administratif') ORDER BY id DESC",(sid,))
    return render_template('school_accounts.html',school=school,accounts=accounts,matricule_rules={r['role']:dict(r) for r in allrows('SELECT * FROM matricule_rules WHERE school_id=?',(sid,))},base_url=os.getenv('PUBLIC_BASE_URL',request.host_url.rstrip('/')).rstrip('/'))

@app.route('/schools/<int:sid>/accounts/<int:uid>/edit',methods=['GET','POST'])
@roles('superadmin')
def school_account_edit(sid,uid):
    school=one('SELECT id,name,slug FROM schools WHERE id=?',(sid,))
    account=one("SELECT * FROM users WHERE id=? AND school_id=? AND role IN ('admin','administratif')",(uid,sid))
    if not school or not account:abort(404)
    if request.method=='POST':
        first=request.form.get('first_name','').strip()
        last=request.form.get('last_name','').strip()
        email=request.form.get('email','').strip().lower()
        phone=request.form.get('phone','').strip()
        role=request.form.get('role','')
        if not first or not last or not phone or not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+',email):abort(400)
        if len(first)>90 or len(last)>90 or len(phone)>40 or len(email)>254 or role not in ('admin','administratif'):abort(400)
        duplicate=one('SELECT id FROM users WHERE school_id=? AND email=? AND id<>?',(sid,email,uid))
        if duplicate:
            flash('Adresse e-mail déjà utilisée.','error')
            return redirect(url_for('school_account_edit',sid=sid,uid=uid))
        try:
            with db() as c:
                c.execute('BEGIN IMMEDIATE')
                matricule=validate_matricule(request.form.get('matricule') or account['matricule']) if (request.form.get('matricule') or account['matricule']) else next_matricule(c,sid,role)
                if c.execute('SELECT 1 FROM users WHERE school_id=? AND matricule=? AND id<>?',(sid,matricule,uid)).fetchone():
                    raise ValueError('Matricule déjà utilisé dans cet établissement')
                c.execute('UPDATE users SET name=?,first_name=?,last_name=?,email=?,phone=?,role=?,matricule=? WHERE id=? AND school_id=?',
                          (first+' '+last,first,last,email,phone,role,matricule,uid,sid))
        except (ValueError,sqlite3.IntegrityError) as exc:
            flash(str(exc),'error')
            return redirect(url_for('school_account_edit',sid=sid,uid=uid))
        with db() as c:
            c.execute('INSERT INTO audit(school_id,user_id,action,details) VALUES (?,?,?,?)',
                      (sid,current_user()['id'],'account.updated',str(uid)))
        flash('Compte modifié avec succès.','success')
        return redirect(url_for('school_accounts',sid=sid))
    return render_template('school_account_edit.html',school=school,account=account)

@app.post('/schools/<int:sid>/accounts/bulk')
@roles('superadmin')
def school_accounts_bulk(sid):
    school=one('SELECT id,name FROM schools WHERE id=?',(sid,))
    if not school:abort(404)
    values=request.form.getlist('account_ids')
    action=request.form.get('action','')
    if action not in ('suspend','activate','delete','set_admin','set_administratif') or not values or len(values)>100 or any(not v.isdigit() for v in values):
        flash('Sélectionnez des comptes et une action.','error')
        return redirect(url_for('school_accounts',sid=sid))
    ids=list(dict.fromkeys(int(v) for v in values))
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        placeholders=','.join('?' for _ in ids)
        accounts=c.execute(f"SELECT id FROM users WHERE school_id=? AND role IN ('admin','administratif') AND id IN ({placeholders})",(sid,*ids)).fetchall()
        if len(accounts)!=len(ids):abort(403)
        if action=='delete':
            # Historical accounts must be suspended, not deleted: protect academic data and traceability.
            dependencies=(
                ('enrollments','student_id'),('parent_links','student_id'),('parent_links','parent_id'),
                ('courses','teacher_id'),('lessons','teacher_id'),('teacher_history','teacher_id'),
                ('grades','student_id'),('grade_history','actor_id'),
                ('attendance','student_id'),('submissions','student_id'),('unlocks','teacher_id'),
            )
            for uid in ids:
                for table,col in dependencies:
                    if c.execute(f'SELECT 1 FROM {table} WHERE {col}=? LIMIT 1',(uid,)).fetchone():
                        flash('Suppression refusée : certains comptes sont liés à des données scolaires. Vous pouvez les suspendre.','error')
                        return redirect(url_for('school_accounts',sid=sid))
            for uid in ids:
                c.execute('DELETE FROM invites WHERE user_id=?',(uid,))
                c.execute('DELETE FROM notifications WHERE user_id=?',(uid,))
                c.execute('DELETE FROM push_subscriptions WHERE user_id=?',(uid,))
                c.execute('UPDATE audit SET user_id=NULL WHERE user_id=?',(uid,))
            c.execute(f"DELETE FROM users WHERE school_id=? AND id IN ({placeholders})",(sid,*ids))
        elif action in ('activate','suspend'):
            c.execute(f"UPDATE users SET active=? WHERE school_id=? AND id IN ({placeholders})",
                      (1 if action=='activate' else 0,sid,*ids))
        else:
            role='admin' if action=='set_admin' else 'administratif'
            c.execute(f"UPDATE users SET role=? WHERE school_id=? AND id IN ({placeholders})",(role,sid,*ids))
        c.execute('INSERT INTO audit(school_id,user_id,action,details) VALUES (?,?,?,?)',
                  (sid,current_user()['id'],'accounts.'+action,','.join(str(v) for v in ids)))
    flash(f'{len(ids)} compte(s) : modification enregistrée.','success')
    return redirect(url_for('school_accounts',sid=sid))

@app.post('/schools/<int:sid>/update')
@roles('superadmin')
def school_update(sid):
    school=one('SELECT * FROM schools WHERE id=?',(sid,))
    if not school:abort(404)
    action=request.form.get('action','update')
    if action=='delete':
        if request.form.get('confirm_slug')!=school['slug']:abort(400,'Confirmation requise.')
        with db() as c:
            c.execute('BEGIN IMMEDIATE')
            c.execute('DELETE FROM attendance WHERE lesson_id IN (SELECT id FROM lessons WHERE school_id=?)',(sid,))
            c.execute('DELETE FROM submissions WHERE assessment_id IN (SELECT id FROM assessments WHERE school_id=?)',(sid,))
            c.execute('DELETE FROM grade_history WHERE grade_id IN (SELECT g.id FROM grades g JOIN assessments a ON a.id=g.assessment_id WHERE a.school_id=?)',(sid,))
            c.execute('DELETE FROM grades WHERE assessment_id IN (SELECT id FROM assessments WHERE school_id=?)',(sid,))
            c.execute('DELETE FROM assessments WHERE school_id=?',(sid,))
            c.execute('DELETE FROM unlocks WHERE period_id IN (SELECT id FROM periods WHERE school_id=?)',(sid,))
            c.execute('DELETE FROM periods WHERE school_id=?',(sid,))
            c.execute('DELETE FROM lessons WHERE school_id=?',(sid,))
            c.execute('DELETE FROM teacher_history WHERE course_id IN (SELECT id FROM courses WHERE school_id=?)',(sid,))
            c.execute('DELETE FROM courses WHERE school_id=?',(sid,))
            c.execute('DELETE FROM bulletin_publications WHERE school_id=?',(sid,))
            c.execute('DELETE FROM enrollments WHERE school_id=?',(sid,))
            c.execute('DELETE FROM parent_links WHERE school_id=?',(sid,))
            c.execute('DELETE FROM classes WHERE school_id=?',(sid,))
            c.execute('DELETE FROM years WHERE school_id=?',(sid,))
            c.execute('DELETE FROM campuses WHERE school_id=?',(sid,))
            c.execute('DELETE FROM notifications WHERE school_id=?',(sid,))
            c.execute('DELETE FROM push_subscriptions WHERE user_id IN (SELECT id FROM users WHERE school_id=?)',(sid,))
            c.execute('DELETE FROM audit WHERE school_id=?',(sid,))
            c.execute('DELETE FROM invites WHERE user_id IN (SELECT id FROM users WHERE school_id=?)',(sid,))
            c.execute('DELETE FROM users WHERE school_id=?',(sid,))
            c.execute('DELETE FROM matricule_rules WHERE school_id=?',(sid,))
            c.execute('DELETE FROM schools WHERE id=?',(sid,))
        flash('Établissement et données associés supprimés définitivement.','success')
    elif action in ('activate','deactivate'):
        execute('UPDATE schools SET active=? WHERE id=?',(1 if action=='activate' else 0,sid))
        flash('Statut de l’établissement mis à jour.','success')
    elif action=='update':
        end=request.form.get('subscription_end','').strip() or None
        if end:
            try:datetime.date.fromisoformat(end)
            except ValueError:abort(400)
        plan=request.form.get('plan',school['plan'])
        if plan not in ('essentiel','standard','premium'):abort(400)
        logo,mime=read_logo(request.files.get('logo'))
        campus_enabled=int(request.form.get('campus_enabled')=='1')
        execute('UPDATE schools SET campus_enabled=? WHERE id=?',(campus_enabled,sid))
        if logo:execute('UPDATE schools SET plan=?,subscription_end=?,logo=?,logo_mime=? WHERE id=?',(plan,end,logo,mime,sid))
        elif request.form.get('remove_logo')=='1':execute('UPDATE schools SET plan=?,subscription_end=?,logo=NULL,logo_mime=NULL WHERE id=?',(plan,end,sid))
        else:execute('UPDATE schools SET plan=?,subscription_end=? WHERE id=?',(plan,end,sid))
        flash('Paramètres de l’établissement enregistrés.','success')
    else:abort(400)
    return redirect(url_for('schools'))

@app.route('/settings',methods=['GET','POST'])
@roles('superadmin')
def settings():
    if request.method=='POST':
        name=request.form.get('support_name','').strip()
        phone=request.form.get('support_phone','').strip()
        powered=request.form.get('powered_by','').strip()
        if not name or not phone or not powered or len(name)>120 or len(phone)>40 or len(powered)>180:
            flash('Vérifiez les champs du support : nom (120 caractères), téléphone (40) et signature (180) sont obligatoires.','error')
            return redirect(url_for('settings'))
        try:
            logo,mime=read_logo(request.files.get('platform_logo'))
        except HTTPException as exc:
            flash(exc.description or 'Fichier logo invalide. Formats autorisés : PNG, JPG ou WebP (5 Mo maximum).','error')
            return redirect(url_for('settings'))
        remove=request.form.get('remove_platform_logo')=='1'
        if logo and remove:
            flash('Choisissez soit un nouveau logo, soit sa suppression.','error')
            return redirect(url_for('settings'))
        with db() as c:
            if logo:
                c.execute('INSERT INTO platform_branding(id,logo,logo_mime) VALUES (1,?,?) ON CONFLICT(id) DO UPDATE SET logo=excluded.logo,logo_mime=excluded.logo_mime',(logo,mime))
            elif remove:
                c.execute('DELETE FROM platform_branding WHERE id=1')
            for k,v in [('support_name',name),('support_phone',phone),('powered_by',powered)]:
                c.execute('INSERT INTO settings(key,value) VALUES (?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value',(k,v))
        flash('Paramètres de la plateforme et identité visuelle enregistrés.','success')
        return redirect(url_for('settings'))
    return render_template('settings.html')

def assign_courses_to_teacher(teacher_id, course_ids):
    """Assign several existing course offerings, retaining teacher history."""
    wanted=list(dict.fromkeys(int(x) for x in course_ids if str(x).isdigit()))
    if len(wanted)>100:abort(400,'Trop de cours sélectionnés')
    if not wanted:return
    marker=','.join('?' for _ in wanted)
    matches=allrows(f'SELECT id,teacher_id FROM courses WHERE school_id=? AND id IN ({marker})',(school_id(),*wanted))
    if len(matches)!=len(wanted):abort(403,'Cours en dehors de votre établissement')
    now=datetime.datetime.now(datetime.timezone.utc).isoformat()
    with db() as c:
        for course in matches:
            if course['teacher_id']==teacher_id:continue
            if course['teacher_id']:
                c.execute('UPDATE teacher_history SET date_end=? WHERE course_id=? AND date_end IS NULL',(now,course['id']))
            c.execute('UPDATE courses SET teacher_id=? WHERE id=?',(teacher_id,course['id']))
            c.execute('INSERT INTO teacher_history(course_id,teacher_id,date_start) VALUES (?,?,?)',(course['id'],teacher_id,now))


@app.get('/api/matricule/next')
@roles('admin','administratif')
def matricule_preview():
    role=request.args.get('role','etudiant')
    if role not in MATRICULE_PROFILES:abort(400)
    with db() as c:
        value=next_matricule(c,school_id(),role,advance=False)
    return jsonify(matricule=value)

@app.get('/api/schools/<int:sid>/matricule/next')
@roles('superadmin')
def school_matricule_preview(sid):
    if not one('SELECT id FROM schools WHERE id=?',(sid,)):abort(404)
    role=request.args.get('role','admin')
    if role not in MATRICULE_PROFILES:abort(400)
    with db() as c:
        value=next_matricule(c,sid,role,advance=False)
    return jsonify(matricule=value)

@app.route('/schools/<int:sid>/matricules',methods=['GET','POST'])
@roles('superadmin')
def school_matricules(sid):
    school=one('SELECT id,name FROM schools WHERE id=?',(sid,))
    if not school:abort(404)
    if request.method=='POST':
        try:
            with db() as c:
                c.execute('BEGIN IMMEDIATE')
                ensure_matricule_rules(c,sid)
                for role in MATRICULE_PROFILES:
                    prefix=request.form.get('prefix_'+role,'').strip().upper()
                    padding=int(request.form.get('padding_'+role,'4'))
                    start=int(request.form.get('next_'+role,'1'))
                    if not re.fullmatch(r'[A-Z0-9][A-Z0-9_/-]{0,14}',prefix) or not (2<=padding<=10) or not (1<=start<=999999999):
                        raise ValueError('Paramètres invalides pour '+role)
                    c.execute('UPDATE matricule_rules SET prefix=?,padding=?,next_number=? WHERE school_id=? AND role=?',
                              (prefix,padding,start,sid,role))
            flash('Numérotation des matricules enregistrée.','success')
        except (ValueError,sqlite3.IntegrityError) as exc:
            flash(str(exc),'error')
        return redirect(url_for('school_matricules',sid=sid))
    with db() as c:
        ensure_matricule_rules(c,sid)
    rules=allrows('SELECT * FROM matricule_rules WHERE school_id=? ORDER BY role',(sid,))
    return render_template('school_matricules.html',school=school,rules=rules)

@app.post('/accounts/<int:uid>/resend-invite')
@roles('superadmin','admin','administratif')
def resend_account_invite(uid):
    actor=current_user()
    person=one('SELECT id,school_id,role FROM users WHERE id=? AND active=1',(uid,))
    if not person:abort(404)
    if actor['role']=='superadmin':
        if person['role'] not in ('admin','administratif'):abort(403)
        target=url_for('school_accounts',sid=person['school_id'])
    else:
        if person['school_id']!=actor['school_id'] or person['role']=='superadmin':abort(403)
        target=url_for('people')
    sent=create_invite(uid)
    flash('Lien envoyé par e-mail.' if sent else 'Brevo indisponible : invitation non envoyée.', 'success' if sent else 'error')
    return redirect(target)

@app.post('/people/<int:uid>/matricule')
@roles('admin')
def edit_person_matricule(uid):
    user=one('SELECT id,matricule FROM users WHERE school_id=? AND id=?',(school_id(),uid))
    if not user:abort(404)
    try:
        value=validate_matricule(request.form.get('matricule'))
        with db() as c:
            c.execute('BEGIN IMMEDIATE')
            if c.execute('SELECT 1 FROM users WHERE school_id=? AND matricule=? AND id<>?',(school_id(),value,uid)).fetchone():
                raise ValueError('Matricule déjà attribué')
            c.execute('UPDATE users SET matricule=? WHERE school_id=? AND id=?',(value,school_id(),uid))
        audit('person.matricule_updated',f'{uid}:{user["matricule"]}->{value}')
        flash('Matricule mis à jour.','success')
    except (ValueError,sqlite3.IntegrityError) as exc:
        flash(str(exc),'error')
    return redirect(url_for('people'))

@app.route('/people',methods=['GET','POST'])
@roles('admin','administratif')
def people():
    if request.method=='POST':
        role=request.form.get('role','etudiant')
        if role not in ('etudiant','enseignant','parent','administratif'):abort(400)
        name=request.form.get('name','').strip()
        email=request.form.get('email','').strip().lower()
        if not name or not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+',email):abort(400)
        class_id=request.form.get('class_id','').strip()
        if role=='etudiant':
            if not class_id or not has_school('classes',class_id):
                flash('Sélectionnez une classe de votre établissement.','error')
                return redirect(url_for('people'))
        elif class_id:abort(400,'Seuls les étudiants peuvent être inscrits dans une classe depuis ce formulaire')
        course_ids=request.form.getlist('course_ids') if role=='enseignant' else []
        if course_ids:
            if any(not str(v).isdigit() for v in course_ids):abort(400)
            vals=list(dict.fromkeys(int(v) for v in course_ids))
            clause=','.join('?' for _ in vals)
            selected=allrows(f'SELECT id FROM courses WHERE school_id=? AND id IN ({clause})',(school_id(),*vals))
            if len(selected)!=len(vals):abort(403)
        password=secrets.token_urlsafe(32)
        with db() as c:
            if c.execute('SELECT id FROM users WHERE school_id=? AND email=?',(school_id(),email)).fetchone():
                flash('Adresse e-mail déjà utilisée.','error')
                return redirect(url_for('people'))
            c.execute('BEGIN IMMEDIATE')
            try:
                matricule=requested_matricule(c,school_id(),role,request.form.get('matricule'))
                uid=c.execute('INSERT INTO users(school_id,name,email,password_hash,role,matricule,must_change_password) VALUES (?,?,?,?,?,?,1)',(school_id(),name,email,generate_password_hash(password),role,matricule)).lastrowid
            except (ValueError,sqlite3.IntegrityError) as exc:
                flash(str(exc),'error')
                return redirect(url_for('people'))
            if role=='etudiant':
                clazz=c.execute('SELECT year_id FROM classes WHERE school_id=? AND id=?',(school_id(),class_id)).fetchone()
                if not clazz:abort(403)
                c.execute('INSERT INTO enrollments(school_id,student_id,class_id,year_id) VALUES (?,?,?,?)',(school_id(),uid,int(class_id),clazz['year_id']))
        if course_ids:assign_courses_to_teacher(uid,course_ids)
        audit('person.created',email)
        sent=create_invite(uid)
        flash('Compte créé. '+('Lien d’activation envoyé par e-mail.' if sent else 'Invitation non envoyée : vérifiez la configuration Brevo.'),'success' if sent else 'error')
        return redirect(url_for('people'))
    users=scoped('SELECT * FROM users WHERE school_id=? ORDER BY role,name')
    classes=scoped('SELECT id,name,year_id FROM classes WHERE school_id=? ORDER BY name')
    years=scoped('SELECT id,label FROM years WHERE school_id=? ORDER BY label DESC')
    courses=scoped('SELECT c.id,c.name,cl.name class_name FROM courses c JOIN classes cl ON cl.id=c.class_id WHERE c.school_id=? ORDER BY cl.name,c.name')
    rows=allrows("""SELECT u.id,u.name,u.email,u.role,u.matricule,u.active,u.must_change_password,
                     GROUP_CONCAT(DISTINCT cl.name) class_names,
                     TRIM(COALESCE(GROUP_CONCAT(DISTINCT cl.id),'') || ',' ||
                       COALESCE((SELECT GROUP_CONCAT(DISTINCT c2.class_id) FROM courses c2 WHERE c2.school_id=u.school_id AND c2.teacher_id=u.id),''), ',') class_ids,
                     TRIM(COALESCE(GROUP_CONCAT(DISTINCT e.year_id),'') || ',' ||
                       COALESCE((SELECT GROUP_CONCAT(DISTINCT cl2.year_id) FROM courses c2 JOIN classes cl2 ON cl2.id=c2.class_id WHERE c2.school_id=u.school_id AND c2.teacher_id=u.id),''), ',') year_ids,
                     (SELECT GROUP_CONCAT(c.name||' ('||cl2.name||')',', ') FROM courses c JOIN classes cl2 ON cl2.id=c.class_id WHERE c.teacher_id=u.id AND c.school_id=u.school_id) teacher_courses
                  FROM users u LEFT JOIN enrollments e ON e.student_id=u.id AND e.school_id=u.school_id
                  LEFT JOIN classes cl ON cl.id=e.class_id
                  WHERE u.school_id=? GROUP BY u.id ORDER BY u.role,u.name""",(school_id(),))
    return render_template('list.html',title='Utilisateurs',columns=['Matricule','Nom','E-mail','Profil','Classe / cours attribués'],rows=[(x['matricule'],x['name'],x['email'],x['role'],x['class_names'] if x['role']=='etudiant' else x['teacher_courses'] or '—') for x in rows],form='people',classes=classes,years=years,courses=courses,user_rows=rows,matricule_rules={r['role']:dict(r) for r in allrows('SELECT * FROM matricule_rules WHERE school_id=?',(school_id(),))})

@app.post('/people/bulk')
@roles('admin','administratif')
def people_bulk():
    values=request.form.getlist('user_ids')
    action=request.form.get('action','')
    if action not in ('suspend','activate','delete','invite') or not values or len(values)>200 or any(not x.isdigit() for x in values):
        flash('Sélectionnez des utilisateurs et une action valide.','error')
        return redirect(url_for('people'))
    selected=list(dict.fromkeys(int(x) for x in values))
    if action=='invite' and len(selected)>30:
        flash('Limitez les envois d’invitation à 30 utilisateurs par opération.','error')
        return redirect(url_for('people'))
    if action=='invite':
        matches=allrows('SELECT id,role FROM users WHERE school_id=? AND active=1 AND id IN ('+','.join('?' for _ in selected)+')',(school_id(),*selected))
        if len(matches)!=len(selected) or current_user()['id'] in selected or any(x['role']=='superadmin' for x in matches):abort(403)
        successes=sum(1 for person in matches if create_invite(person['id']))
        flash(f'{successes} invitation(s) envoyée(s) sur {len(matches)}.','success' if successes==len(matches) else 'error')
        return redirect(url_for('people'))
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        users=c.execute('SELECT id,role,name FROM users WHERE school_id=? AND id IN ('+','.join('?' for _ in selected)+')',(school_id(),*selected)).fetchall()
        if len(users)!=len(selected) or current_user()['id'] in selected or any(x['role']=='superadmin' for x in users):
            abort(403)
        if action=='delete':
            # Keep grades, attendance, teacher assignments and audited academic history intact.
            for user in users:
                uid=user['id']
                linked=any(c.execute('SELECT 1 FROM '+sql+' LIMIT 1',(uid,)).fetchone() for sql in [
                    'grades WHERE student_id=?','attendance WHERE student_id=?','submissions WHERE student_id=?',
                    'courses WHERE teacher_id=?','lessons WHERE teacher_id=?','teacher_history WHERE teacher_id=?',
                    'grade_history WHERE actor_id=?','unlocks WHERE teacher_id=?'])
                if linked:
                    flash('Suppression refusée : un compte possède des notes, séances ou affectations historisées. Suspendez-le pour préserver les dossiers.','error')
                    return redirect(url_for('people'))
            for user in users:
                uid=user['id']
                c.execute('DELETE FROM parent_links WHERE student_id=? OR parent_id=?',(uid,uid))
                c.execute('DELETE FROM enrollments WHERE student_id=?',(uid,))
                c.execute('DELETE FROM notifications WHERE user_id=?',(uid,))
                c.execute('DELETE FROM invites WHERE user_id=?',(uid,))
                c.execute('UPDATE audit SET user_id=NULL WHERE user_id=?',(uid,))
                c.execute('DELETE FROM users WHERE id=? AND school_id=?',(uid,school_id()))
        else:
            c.execute('UPDATE users SET active=? WHERE school_id=? AND id IN ('+','.join('?' for _ in selected)+')',(1 if action=='activate' else 0,school_id(),*selected))
    audit('users.'+action,','.join(str(x) for x in selected))
    flash(f'{len(selected)} utilisateur(s) : action « {action} » appliquée.','success')
    return redirect(url_for('people'))

@app.get('/people/import/template')
@roles('admin','administratif')
def students_csv_template():
    csv_data='\ufeffprenom;nom;email;classe\nAmadou;Traore;amadou.exemple@example.test;Licence 1\n'
    return Response(csv_data,mimetype='text/csv; charset=utf-8',headers={'Content-Disposition':'attachment; filename=edupilot-modele-etudiants.csv'})

@app.post('/people/import')
@roles('admin','administratif')
def people_import():
    uploaded=request.files.get('file')
    if not uploaded or not uploaded.filename:abort(400,'Fichier obligatoire')
    filename=uploaded.filename.lower()
    if not filename.endswith(('.csv','.xlsx')):abort(400,'Format CSV ou XLSX uniquement')
    content=uploaded.read(2*1024*1024+1)
    if len(content)>2*1024*1024:abort(400,'Limite de 2 Mo')
    try:
        if filename.endswith('.csv'):
            text=content.decode('utf-8-sig')
            dialect=csv.Sniffer().sniff(text[:2048],delimiters=',;\t') if text.strip() else csv.excel
            rows=list(csv.reader(io.StringIO(text),dialect))
        else:
            workbook=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
            sheet=workbook.active
            rows=list(sheet.iter_rows(values_only=True,max_row=2002))
            workbook.close()
    except Exception as exc:
        app.logger.info('Invalid student import file: %s',type(exc).__name__)
        flash('Impossible de lire le fichier : utilisez un CSV UTF-8 ou un Excel XLSX valide.','error')
        return redirect(url_for('people'))
    if len(rows)<2:
        flash('Le fichier doit contenir un en-tête et au moins un étudiant.','error')
        return redirect(url_for('people'))
    if len(rows)>2001:
        flash('Maximum 2 000 étudiants par import.','error')
        return redirect(url_for('people'))
    def key(s):
        return slugify(str(s or '')).replace('_','')
    headers=[key(x) for x in rows[0]]
    def find_col(*names):
        for name in names:
            if key(name) in headers:return headers.index(key(name))
        return None
    fi=find_col('prenom','first_name')
    la=find_col('nom','last_name')
    em=find_col('email','e_mail','mail')
    cl=find_col('classe','class','class_id')
    selected_class_id=request.form.get('import_class_id','').strip()
    if None in (fi,la,em) or (cl is None and not selected_class_id):
        flash('Colonnes requises : prenom, nom, email. Indiquez une classe dans le formulaire ou dans le fichier.','error')
        return redirect(url_for('people'))
    classes=allrows('SELECT id,name,year_id FROM classes WHERE school_id=?',(school_id(),))
    selected_class=None
    if selected_class_id:
        selected_class=next((c for c in classes if str(c['id'])==selected_class_id),None)
        if selected_class is None:
            flash('Classe sélectionnée introuvable dans votre établissement.','error')
            return redirect(url_for('people'))
    by_id={str(c['id']):c for c in classes}
    by_name={}
    for c in classes:by_name.setdefault(key(c['name']),[]).append(c)
    prepared=[]; errors=[]; encountered=set()
    for lineno,row in enumerate(rows[1:],2):
        if not row or all(v is None or not str(v).strip() for v in row):continue
        def cell(i):return str(row[i] or '').strip() if i<len(row) else ''
        first,last,email=cell(fi),cell(la),cell(em).lower()
        lookup=cell(cl) if cl is not None else ''
        target=selected_class
        if target is None:
            target=by_id.get(lookup)
            if not target:
                opts=by_name.get(key(lookup),[])
                target=opts[0] if len(opts)==1 else None
        if not first or not last or not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+',email):
            errors.append(f'Ligne {lineno} : identité ou e-mail invalide')
        elif not target:
            errors.append(f'Ligne {lineno} : classe introuvable ou ambiguë ({lookup})')
        elif email in encountered:
            errors.append(f'Ligne {lineno} : e-mail en doublon ({email})')
        else:
            encountered.add(email)
            prepared.append((first+' '+last,email,target['id'],target['year_id']))
        if len(errors)>=8:break
    if not prepared and not errors:errors.append('Aucune ligne exploitable')
    if errors:
        flash('Import annulé. '+ ' ; '.join(errors),'error')
        return redirect(url_for('people'))
    hashes=[]
    try:
        with db() as c:
            c.execute('BEGIN IMMEDIATE')
            registered={r['email'] for r in c.execute('SELECT email FROM users WHERE school_id=? AND email IN ('+','.join('?' for _ in prepared)+')',(school_id(),*tuple(x[1] for x in prepared))).fetchall()}
            if registered:
                flash('Import annulé : comptes déjà existants : '+', '.join(sorted(registered)[:4]),'error')
                return redirect(url_for('people'))
            for name,email,class_id,year_id in prepared:
                uid=c.execute('INSERT INTO users(school_id,name,email,password_hash,role,matricule) VALUES (?,?,?,?,?,?)',(school_id(),name,email,generate_password_hash(secrets.token_urlsafe(24)),'etudiant',next_matricule(c,school_id(),'etudiant'))).lastrowid
                c.execute('INSERT INTO enrollments(school_id,student_id,class_id,year_id) VALUES (?,?,?,?)',(school_id(),uid,class_id,year_id))
                hashes.append(uid)
    except sqlite3.IntegrityError:
        flash('Import annulé : conflit sur les données.','error')
        return redirect(url_for('people'))
    # Invitations can be requested later; avoid thousands of synchronous outbound email calls.
    audit('students.imported',str(len(hashes)))
    flash(f'{len(hashes)} étudiants importés et inscrits dans '+(selected_class['name'] if selected_class else 'leurs classes')+'. Invitations disponibles via l’action d’envoi groupé.','success')
    return redirect(url_for('people'))

def create_invite(uid):
    return send_account_link(uid,'activation')

@app.route('/activate/<token>',methods=['GET','POST'])
def activate(token):
    hashed=hashlib.sha256(token.encode()).hexdigest()
    invitation=one('SELECT * FROM invites WHERE token_hash=? AND used=0',(hashed,))
    now=datetime.datetime.now(datetime.timezone.utc).isoformat()
    if not invitation or invitation['expires_at']<now:abort(404,'Lien expiré ou déjà utilisé.')
    if request.method=='POST':
        password=request.form.get('password','')
        confirm=request.form.get('confirm_password','')
        if not valid_password(password):
            flash('Au moins 12 caractères, avec lettres et chiffres.','error')
        elif password!=confirm:
            flash('Les deux mots de passe ne correspondent pas.','error')
        else:
            with db() as c:
                c.execute('BEGIN IMMEDIATE')
                claim=c.execute('UPDATE invites SET used=1 WHERE id=? AND used=0 AND expires_at>=?',(invitation['id'],datetime.datetime.now(datetime.timezone.utc).isoformat()))
                if claim.rowcount!=1:abort(404,'Lien expiré ou déjà utilisé.')
                c.execute('UPDATE users SET password_hash=?,must_change_password=0 WHERE id=?',(generate_password_hash(password),invitation['user_id']))
                c.execute('DELETE FROM invites WHERE user_id=? AND id<>?',(invitation['user_id'],invitation['id']))
            flash('Mot de passe défini. Connectez-vous à EduPilot.','success')
            return redirect(url_for('login'))
    return render_template('activate.html')

def valid_password(value):
    return len(value)>=12 and any(c.isalpha() for c in value) and any(c.isdigit() for c in value)

@app.route('/forgot-password',methods=['GET','POST'])
def forgot_password():
    if request.method=='POST':
        email=request.form.get('email','').strip().lower()
        if len(email)>254:abort(400)
        # Always answer identically, regardless of whether the address exists.
        # Throttle per address and source IP; never store the address in the rate-limit table.
        ip=(request.remote_addr or 'unknown')[:80]
        stamp=int(time.time())
        keys=[hashlib.sha256((key+':'+value).encode()).hexdigest()
              for key,value in [('email',email),('ip',ip)]]
        permitted=False
        with db() as c:
            c.execute('BEGIN IMMEDIATE')
            recent=[c.execute('SELECT attempted_at FROM password_reset_limits WHERE key_hash=?',(key,)).fetchone() for key in keys]
            if all(not row or stamp-row['attempted_at']>=60 for row in recent):
                permitted=True
                for key in keys:
                    c.execute('INSERT INTO password_reset_limits(key_hash,attempted_at) VALUES (?,?) ON CONFLICT(key_hash) DO UPDATE SET attempted_at=excluded.attempted_at',(key,stamp))
        users=allrows('SELECT id FROM users WHERE email=? AND active=1',(email,)) if permitted else []
        # All accounts with this verified mailbox receive their own distinct reset link.
        for user in users:send_account_link(user['id'],'reset')
        flash('Si cette adresse correspond à un compte actif, un lien de réinitialisation vient d’être envoyé.','success')
        return redirect(url_for('forgot_password'))
    return render_template('forgot_password.html')

@app.route('/classes',methods=['GET','POST'])
@roles('admin','administratif')
def classes():
    if request.method=='POST':
        year=scoped('SELECT id FROM years WHERE school_id=? AND is_active=1 LIMIT 1')
        if not year:abort(400)
        campus=request.form.get('campus_id')
        if campus and (not one('SELECT campus_enabled FROM schools WHERE id=?',(school_id(),))['campus_enabled'] or not has_school('campuses',campus)):abort(403)
        if campus and one('SELECT plan FROM schools WHERE id=?',(school_id(),))['plan']!='premium':
            first=one('SELECT id FROM campuses WHERE school_id=? ORDER BY id LIMIT 1',(school_id(),))
            if int(campus)!=first['id']: abort(403)
        name=request.form.get('name','').strip()
        if not name:abort(400)
        cid=execute('INSERT INTO classes(school_id,year_id,campus_id,name,cycle,parents_enabled) VALUES (?,?,?,?,?,?)',(school_id(),year[0]['id'],campus or None,name,request.form.get('cycle','semestre'),int('parents_enabled' in request.form)))
        execute('INSERT INTO periods(school_id,class_id,name) VALUES (?,?,?)',(school_id(),cid,'Trimestre 1' if request.form.get('cycle')=='trimestre' else 'Semestre 1'))
        audit('class.created',name)
        return redirect(url_for('classes'))
    return render_template('list.html',title='Classes',columns=['Classe','Régime','Parents','Détails','Modifier'],rows=[(x['name'],x['cycle'],'Actif' if x['parents_enabled'] else 'Inactif',{'link':url_for('class_details',cid=x['id']),'label':'Détails'},{'link':url_for('class_edit',cid=x['id']),'label':'Modifier'}) for x in scoped('SELECT * FROM classes WHERE school_id=?')],form='classes',campuses=scoped('SELECT * FROM campuses WHERE school_id=?') if one('SELECT campus_enabled FROM schools WHERE id=?',(school_id(),))['campus_enabled'] else [])

@app.route('/classes/<int:cid>/edit',methods=['GET','POST'])
@roles('admin','administratif')
def class_edit(cid):
    item=one('SELECT * FROM classes WHERE id=? AND school_id=?',(cid,school_id()))
    if not item:abort(404)
    if request.method=='POST':
        name=request.form.get('name','').strip()
        cycle=request.form.get('cycle','')
        campus=request.form.get('campus_id') or None
        if not name or len(name)>180 or cycle not in ('semestre','trimestre'):abort(400)
        enabled=one('SELECT campus_enabled FROM schools WHERE id=?',(school_id(),))['campus_enabled']
        if not enabled:campus=None
        if campus and not has_school('campuses',campus):abort(403)
        execute('UPDATE classes SET name=?,cycle=?,campus_id=?,parents_enabled=? WHERE id=? AND school_id=?',(name,cycle,campus,int('parents_enabled' in request.form),cid,school_id()))
        audit('class.updated',str(cid));flash('Classe modifiée.','success')
        return redirect(url_for('classes'))
    return render_template('edit_academic.html',kind='class',item=item,campuses=scoped('SELECT * FROM campuses WHERE school_id=?') if one('SELECT campus_enabled FROM schools WHERE id=?',(school_id(),))['campus_enabled'] else [])

@app.get('/classes/<int:cid>/details')
@roles('admin','administratif','enseignant')
def class_details(cid):
    clazz=one("""SELECT cl.*, y.label academic_year,ca.name campus_name
                FROM classes cl JOIN years y ON y.id=cl.year_id
                LEFT JOIN campuses ca ON ca.id=cl.campus_id
                WHERE cl.id=? AND cl.school_id=?""",(cid,school_id()))
    if not clazz:abort(404)
    if current_user()['role']=='enseignant' and not one(
        'SELECT id FROM courses WHERE school_id=? AND class_id=? AND teacher_id=?',(school_id(),cid,current_user()['id'])):
        abort(403)
    students=allrows("""SELECT DISTINCT u.id,u.name,u.matricule,u.email,u.active
                        FROM enrollments e JOIN users u ON u.id=e.student_id
                        WHERE e.school_id=? AND e.class_id=? AND e.year_id=?
                        ORDER BY u.name""",(school_id(),cid,clazz['year_id']))
    courses=allrows("""SELECT c.id,c.name,c.code,c.planned_minutes,u.id teacher_id,u.name teacher_name,u.matricule teacher_matricule
                       FROM courses c LEFT JOIN users u ON u.id=c.teacher_id
                       WHERE c.school_id=? AND c.class_id=? ORDER BY c.name""",(school_id(),cid))
    teachers=allrows("""SELECT DISTINCT u.id,u.name,u.matricule,u.email
                        FROM courses c JOIN users u ON u.id=c.teacher_id
                        WHERE c.school_id=? AND c.class_id=? AND u.role IN ('enseignant','admin')
                        ORDER BY u.name""",(school_id(),cid))
    return render_template('class_details.html',clazz=clazz,students=students,
                           teachers=teachers,courses=courses)

@app.post('/enroll')
@roles('admin','administratif')
def enroll():
    student=request.form.get('student_id');cl=request.form.get('class_id')
    if not has_school('classes',cl) or not one("SELECT id FROM users WHERE id=? AND school_id=? AND role='etudiant'",(student,school_id())):abort(403)
    clazz=one('SELECT * FROM classes WHERE id=? AND school_id=?',(cl,school_id()))
    execute('INSERT OR IGNORE INTO enrollments(school_id,student_id,class_id,year_id) VALUES (?,?,?,?)',(school_id(),student,cl,clazz['year_id']))
    audit('student.enrolled',f'{student}:{cl}');flash('Inscription enregistrée.','success')
    return redirect(url_for('students'))

@app.post('/students/transfer')
@roles('admin','administratif')
def students_transfer():
    ids=request.form.getlist('enrollment_ids')
    class_id=request.form.get('class_id')
    dest=one('SELECT id,year_id FROM classes WHERE id=? AND school_id=?',(class_id,school_id()))
    if not dest or not ids or len(ids)>300 or any(not x.isdigit() for x in ids):abort(400)
    ids=list(dict.fromkeys(map(int,ids)))
    now=datetime.datetime.now(datetime.timezone.utc).isoformat()
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        q=','.join('?'*len(ids))
        rows=c.execute(f'SELECT id,student_id,class_id,year_id FROM enrollments WHERE school_id=? AND id IN ({q})',(school_id(),*ids)).fetchall()
        if len(rows)!=len(ids):abort(403)
        # Keep all historical enrollments for past years. Move only records in same academic year.
        if any(r['year_id']!=dest['year_id'] for r in rows):
            flash('Transfert annulé : choisissez une classe de la même année académique. Pour une nouvelle année, utilisez la promotion.','error')
            return redirect(url_for('students'))
        for row in rows:
            if row['class_id']==dest['id']:continue
            duplicate=c.execute('SELECT id FROM enrollments WHERE school_id=? AND student_id=? AND class_id=? AND year_id=?',(school_id(),row['student_id'],dest['id'],dest['year_id'])).fetchone()
            if duplicate:
                c.execute('DELETE FROM enrollments WHERE id=?',(row['id'],))
            else:
                c.execute('UPDATE enrollments SET class_id=? WHERE id=?',(dest['id'],row['id']))
            c.execute('INSERT INTO audit(school_id,user_id,action,details) VALUES (?,?,?,?)',(school_id(),current_user()['id'],'enrollment.transferred',f"{row['student_id']}:{row['class_id']}->{dest['id']}:{now}"))
    flash(f'{len(ids)} inscription(s) transférée(s).','success')
    return redirect(url_for('students'))

@app.route('/students')
@login_required
def students():
    u=current_user()
    if u['role'] not in ('admin','administratif','enseignant'):abort(403)
    sql="""SELECT e.id enrollment_id,u.name,u.matricule,c.name class_name,c.id class_id,y.id year_id,y.label
           FROM enrollments e JOIN users u ON u.id=e.student_id
           JOIN classes c ON c.id=e.class_id JOIN years y ON y.id=e.year_id
           WHERE e.school_id=?"""
    args=[school_id()]
    if u['role']=='enseignant':
        sql+=" AND EXISTS (SELECT 1 FROM courses co WHERE co.class_id=e.class_id AND co.school_id=e.school_id AND co.teacher_id=?)"
        args.append(u['id'])
    requested_class=request.args.get('class_id','')
    requested_year=request.args.get('year_id','')
    if requested_class and requested_class.isdigit():
        sql+=' AND e.class_id=?';args.append(int(requested_class))
    if requested_year and requested_year.isdigit():
        sql+=' AND e.year_id=?';args.append(int(requested_year))
    sql+=' ORDER BY y.label DESC,c.name,u.name'
    entries=allrows(sql,args)
    class_choices=scoped('SELECT id,name,year_id FROM classes WHERE school_id=? ORDER BY name')
    year_choices=scoped('SELECT id,label FROM years WHERE school_id=? ORDER BY label DESC')
    if u['role']=='enseignant':
        class_choices=[c for c in class_choices if one('SELECT id FROM courses WHERE class_id=? AND school_id=? AND teacher_id=?',(c['id'],school_id(),u['id']))]
    return render_template('students.html',enrollments=entries,classes=class_choices,years=year_choices,can_edit=u['role'] in ('admin','administratif'),chosen_class=requested_class,chosen_year=requested_year)

@app.route('/courses',methods=['GET','POST'])
@roles('admin','administratif')
def courses():
    if request.method=='POST':
        cl=request.form.get('class_id');teacher=request.form.get('teacher_id') or None
        if not has_school('classes',cl):abort(403)
        if teacher and not one("SELECT id FROM users WHERE school_id=? AND id=? AND role IN ('enseignant','admin')",(school_id(),teacher)):abort(403)
        name=request.form.get('name','').strip()
        if not name:abort(400)
        try:
            coefficient=float(request.form.get('coefficient','1'))
            if not (0<coefficient<=100):raise ValueError()
        except ValueError:abort(400,'Coefficient matière invalide')
        cid=execute('INSERT INTO courses(school_id,class_id,name,code,planned_minutes,teacher_id,coefficient) VALUES (?,?,?,?,?,?,?)',(school_id(),cl,name,request.form.get('code',''),int(float(request.form.get('hours',20))*60),teacher,coefficient))
        if teacher:execute('INSERT INTO teacher_history(course_id,teacher_id,date_start) VALUES (?,?,?)',(cid,teacher,datetime.date.today().isoformat()))
        audit('course.created',name);return redirect(url_for('courses'))
    rows=scoped('SELECT c.*,cl.name class_name,u.name teacher_name,u.matricule teacher_matricule FROM courses c JOIN classes cl ON cl.id=c.class_id LEFT JOIN users u ON u.id=c.teacher_id WHERE c.school_id=?')
    return render_template('courses_manage.html',courses=rows,classes=scoped('SELECT * FROM classes WHERE school_id=?'),teachers=scoped("SELECT * FROM users WHERE school_id=? AND role IN ('enseignant','admin')"))

@app.post('/courses/bulk')
@roles('admin','administratif')
def courses_bulk():
    ids=request.form.getlist('course_ids')
    action=request.form.get('action','')
    if action not in ('delete','coefficient','teacher') or not ids or len(ids)>100 or any(not x.isdigit() for x in ids):abort(400)
    ids=list(dict.fromkeys(map(int,ids)))
    qs=','.join('?' for _ in ids)
    teacher_id=None;coefficient=None
    if action=='coefficient':
        try:
            coefficient=float(request.form.get('coefficient',''))
            if not (0<coefficient<=100):raise ValueError()
        except (ValueError,TypeError):abort(400,'Coefficient invalide')
    if action=='teacher':
        value=request.form.get('teacher_id','').strip()
        if value:
            if not value.isdigit() or not one("SELECT id FROM users WHERE id=? AND school_id=? AND role IN ('enseignant','admin') AND active=1",(value,school_id())):abort(403)
            teacher_id=int(value)
    try:
        with db() as c:
            c.execute('BEGIN IMMEDIATE')
            rows=c.execute(f'SELECT id,teacher_id FROM courses WHERE school_id=? AND id IN ({qs})',(school_id(),*ids)).fetchall()
            if len(rows)!=len(ids):abort(403)
            if action=='delete':
                for cid in ids:
                    for table in ('lessons','assessments'):
                        if c.execute(f'SELECT id FROM {table} WHERE course_id=? LIMIT 1',(cid,)).fetchone():
                            flash('Suppression refusée : un cours possède des séances ou évaluations. Conservez les historiques.','error')
                            return redirect(url_for('courses'))
                c.execute(f'DELETE FROM teacher_history WHERE course_id IN ({qs})',ids)
                c.execute(f'DELETE FROM courses WHERE school_id=? AND id IN ({qs})',(school_id(),*ids))
            elif action=='coefficient':
                c.execute(f'UPDATE courses SET coefficient=? WHERE school_id=? AND id IN ({qs})',(coefficient,school_id(),*ids))
            else:
                now=datetime.datetime.now(datetime.timezone.utc).isoformat()
                for course in rows:
                    if course['teacher_id']==teacher_id:continue
                    c.execute('UPDATE teacher_history SET date_end=? WHERE course_id=? AND date_end IS NULL',(now,course['id']))
                    if teacher_id:c.execute('INSERT INTO teacher_history(course_id,teacher_id,date_start) VALUES (?,?,?)',(course['id'],teacher_id,now))
                    sync_course_future_teacher(c,course['id'],teacher_id)
                    c.execute('UPDATE courses SET teacher_id=? WHERE id=? AND school_id=?',(teacher_id,course['id'],school_id()))
    except ValueError as exc:
        flash(str(exc),'error')
        return redirect(url_for('courses'))
    audit('courses.bulk.'+action,','.join(map(str,ids)))
    flash(f'{len(ids)} cours : modifications enregistrées.','success')
    return redirect(url_for('courses'))

def sync_course_future_teacher(conn,course_id,teacher_id):
    """Synchroniser uniquement les séances futures, programmées et sans historique.
    Une incompatibilité d'horaire bloque le transfert au lieu de créer un double cours.
    """
    now=datetime.datetime.now().strftime('%Y-%m-%dT%H:%M')
    lessons=conn.execute("""SELECT l.* FROM lessons l WHERE l.course_id=? AND l.starts_at>=?
        AND l.status='programmee' AND COALESCE(l.actual_minutes,0)=0
        AND COALESCE(l.validated,0)=0
        AND NOT EXISTS(SELECT 1 FROM attendance a WHERE a.lesson_id=l.id)
        ORDER BY l.starts_at""",(course_id,now)).fetchall()
    if teacher_id:
        for lesson in lessons:
            conflict=conn.execute("""SELECT id FROM lessons WHERE school_id=? AND teacher_id=?
                AND id<>? AND course_id<>?
                AND status NOT IN ('annulee','reportee')
                AND starts_at<? AND ends_at>? LIMIT 1""",
                (lesson['school_id'],teacher_id,lesson['id'],course_id,lesson['ends_at'],lesson['starts_at'])).fetchone()
            if conflict:
                raise ValueError('Conflit de planning : le nouvel enseignant a déjà une séance prévue. Affectation annulée.')
    for lesson in lessons:
        conn.execute('UPDATE lessons SET teacher_id=? WHERE id=?',(teacher_id,lesson['id']))
    return len(lessons)

# Reprise non destructive des séances à venir dont le cours avait déjà changé
# d'enseignant avant l'installation de ce correctif.
with db() as c:
    c.execute('BEGIN IMMEDIATE')
    stale=c.execute("""SELECT DISTINCT l.course_id,c.teacher_id FROM lessons l
        JOIN courses c ON c.id=l.course_id
        WHERE l.starts_at>=? AND l.status='programmee'
        AND COALESCE(l.actual_minutes,0)=0 AND COALESCE(l.validated,0)=0
        AND l.teacher_id IS NOT c.teacher_id
        AND NOT EXISTS(SELECT 1 FROM attendance a WHERE a.lesson_id=l.id)""",
        (datetime.datetime.now().strftime('%Y-%m-%dT%H:%M'),)).fetchall()
    for item in stale:
        try:
            sync_course_future_teacher(c,item['course_id'],item['teacher_id'])
        except ValueError:
            app.logger.warning('Séances non synchronisées pour le cours %s : conflit de planning',item['course_id'])

@app.route('/courses/<int:cid>/edit',methods=['GET','POST'])
@roles('admin','administratif')
def course_edit(cid):
    item=one('SELECT * FROM courses WHERE id=? AND school_id=?',(cid,school_id()))
    if not item:abort(404)
    if request.method=='POST':
        name=request.form.get('name','').strip()
        class_id=request.form.get('class_id')
        teacher=request.form.get('teacher_id') or None
        if not name or len(name)>180 or not has_school('classes',class_id):abort(400)
        if teacher and not one("SELECT id FROM users WHERE id=? AND school_id=? AND role IN ('enseignant','admin') AND active=1",(teacher,school_id())):abort(403)
        try:
            minutes=int(round(float(request.form.get('hours',''))*60))
            coefficient=float(request.form.get('coefficient','1'))
            if not (0<coefficient<=100):raise ValueError()
            if minutes<30 or minutes>300000:raise ValueError()
        except (ValueError,TypeError):abort(400)
        if int(class_id)!=item['class_id'] and (one('SELECT id FROM lessons WHERE course_id=?',(cid,)) or one('SELECT id FROM assessments WHERE course_id=?',(cid,))):
            flash('Ce cours possède des séances ou évaluations : sa classe ne peut plus être modifiée.','error')
            return redirect(url_for('course_edit',cid=cid))
        now=datetime.datetime.now(datetime.timezone.utc).isoformat()
        changed_teacher=item['teacher_id']!=(int(teacher) if teacher else None)
        try:
            with db() as c:
                c.execute('BEGIN IMMEDIATE')
                moved=sync_course_future_teacher(c,cid,int(teacher) if teacher else None)
                if changed_teacher:
                    c.execute('UPDATE teacher_history SET date_end=? WHERE course_id=? AND date_end IS NULL',(now,cid))
                    if teacher:c.execute('INSERT INTO teacher_history(course_id,teacher_id,date_start) VALUES (?,?,?)',(cid,teacher,now))
                c.execute('UPDATE courses SET name=?,code=?,class_id=?,teacher_id=?,planned_minutes=?,coefficient=? WHERE id=? AND school_id=?',(name,request.form.get('code','').strip(),class_id,teacher,minutes,coefficient,cid,school_id()))
        except ValueError as exc:
            flash(str(exc),'error')
            return redirect(url_for('course_edit',cid=cid))
        if changed_teacher:
            notify_event([cid],'Changement d’enseignant pour le cours '+name+'. '+str(moved)+' séance(s) future(s) réattribuée(s).','/planning',teacher_ids=[item['teacher_id'],int(teacher) if teacher else None])
        audit('course.updated',str(cid));flash(f'Cours mis à jour. {moved} séance(s) future(s) synchronisée(s) ; historique conservé.','success')
        return redirect(url_for('courses'))
    return render_template('edit_academic.html',kind='course',item=item,classes=scoped('SELECT * FROM classes WHERE school_id=?'),teachers=scoped("SELECT * FROM users WHERE school_id=? AND role IN ('enseignant','admin') AND active=1"))

@app.route('/planning',methods=['GET','POST'])
@login_required
def planning():
    u=current_user()
    if u['role']=='superadmin':abort(403)
    if request.method=='POST':
        if u['role'] not in ('admin','administratif'):abort(403)
        course_id=request.form.get('course_id')
        class_id=request.form.get('class_id')
        course=one('SELECT * FROM courses WHERE id=? AND school_id=? AND class_id=?',(course_id,school_id(),class_id))
        if not course:
            flash('Sélectionnez une classe puis un cours appartenant à cette classe.','error')
            return redirect(url_for('planning'))
        if not course['teacher_id']:
            flash('Affectez un enseignant à ce cours avant la planification.','error')
            return redirect(url_for('planning'))
        try:
            day=datetime.date.fromisoformat(request.form.get('lesson_date',''))
            start_time=datetime.time.fromisoformat(request.form.get('start_time',''))
            hours=float(request.form.get('duration_hours',''))
            if hours<0.5 or hours>12 or not (hours*2).is_integer():raise ValueError('durée')
            start_dt=datetime.datetime.combine(day,start_time)
            end_dt=start_dt+datetime.timedelta(minutes=round(hours*60))
            if end_dt.date()!=day:raise ValueError('minuit')
            mode=request.form.get('mode','unique')
            if mode not in ('unique','hebdomadaire'):raise ValueError('mode')
            if mode=='hebdomadaire':
                last=datetime.date.fromisoformat(request.form.get('repeat_until',''))
                if last<day or (last-day).days>365:raise ValueError('récurrence')
                dates=[start_dt+datetime.timedelta(days=7*i) for i in range((last-day).days//7+1)]
            else: dates=[start_dt]
            if len(dates)>53:raise ValueError('limite')
        except (ValueError,TypeError,OverflowError):
            flash('Date, heure, durée ou récurrence invalide. Une séance doit finir le même jour et durer de 0,5 à 12 h.','error')
            return redirect(url_for('planning'))
        room=request.form.get('room','').strip()[:80]
        created=[]
        with db() as c:
            c.execute('BEGIN IMMEDIATE')
            for dt in dates:
                end_at=dt+datetime.timedelta(minutes=round(hours*60))
                start_txt=dt.strftime('%Y-%m-%dT%H:%M')
                end_txt=end_at.strftime('%Y-%m-%dT%H:%M')
                clash=c.execute("""SELECT l.id FROM lessons l JOIN courses c ON c.id=l.course_id
                           WHERE l.school_id=? AND l.status NOT IN ('annulee','reportee')
                           AND l.starts_at < ? AND l.ends_at > ?
                           AND (c.class_id=? OR l.teacher_id=? OR (?<>'' AND l.room=?))
                           LIMIT 1""",(school_id(),end_txt,start_txt,class_id,course['teacher_id'],room,room)).fetchone()
                if clash:
                    c.rollback()
                    flash(f'Conflit le {dt.strftime("%d/%m/%Y")} : classe, enseignant ou salle déjà occupé(e). Aucune séance ajoutée.','error')
                    return redirect(url_for('planning'))
                cur=c.execute('INSERT INTO lessons(school_id,course_id,teacher_id,starts_at,ends_at,room) VALUES (?,?,?,?,?,?)',(school_id(),course_id,course['teacher_id'],start_txt,end_txt,room))
                created.append(cur.lastrowid)
        notify_event([course_id],f'{len(created)} séance(s) programmée(s) : '+course['name']+' — début '+dates[0].strftime('%d/%m/%Y à %H:%M'),'/planning',teacher_ids=[course['teacher_id']])
        audit('lesson.created',','.join(str(x) for x in created))
        flash(f'{len(created)} séance(s) ajoutée(s) au planning.','success')
        return redirect(url_for('planning'))
    query='''SELECT l.*,c.name course_name,cl.name class_name FROM lessons l JOIN courses c ON c.id=l.course_id JOIN classes cl ON cl.id=c.class_id WHERE l.school_id=?'''
    args=[school_id()]
    if u['role']=='enseignant' or (u['role']=='admin' and request.args.get('mine')=='1'):
        query+=' AND l.teacher_id=?';args.append(u['id'])
    elif u['role']=='etudiant':query+=' AND EXISTS (SELECT 1 FROM enrollments e WHERE e.student_id=? AND e.class_id=c.class_id)';args.append(u['id'])
    elif u['role']=='parent':query+=' AND EXISTS (SELECT 1 FROM parent_links p JOIN enrollments e ON e.student_id=p.student_id WHERE p.parent_id=? AND p.active=1 AND e.class_id=c.class_id AND cl.parents_enabled=1)';args.append(u['id'])
    events=allrows(query,args)
    return render_template('planning.html',now_iso=datetime.datetime.now().strftime('%Y-%m-%dT%H:%M'),events=events,courses=scoped('SELECT * FROM courses WHERE school_id=? ORDER BY name') if u['role'] in ('admin','administratif') else [],classes=scoped('SELECT id,name FROM classes WHERE school_id=? ORDER BY name') if u['role'] in ('admin','administratif') else [])

def editable_lesson_row(conn,lid):
    return conn.execute("""SELECT l.*,c.class_id FROM lessons l JOIN courses c ON c.id=l.course_id
        WHERE l.id=? AND l.school_id=?""",(lid,school_id())).fetchone()

def lesson_editable(conn,lesson):
    if lesson['status']!='programmee' or lesson['actual_minutes'] or lesson['validated']:return False
    return not conn.execute('SELECT 1 FROM attendance WHERE lesson_id=? LIMIT 1',(lesson['id'],)).fetchone()

def assert_lesson_schedule_available(conn,planned,editing_ids):
    # Check against existing sessions outside of this batch and against each
    # other; never partially modify a set when one session conflicts.
    ignored=set(editing_ids)
    for item in planned:
        overlapping=conn.execute("""SELECT l.id,c.class_id,l.teacher_id,l.room FROM lessons l
              JOIN courses c ON c.id=l.course_id WHERE l.school_id=?
              AND l.status NOT IN ('annulee','reportee')
              AND l.starts_at<? AND l.ends_at>?""",
              (school_id(),item['end'],item['start'])).fetchall()
        for other in overlapping:
            if other['id'] in ignored:continue
            if other['class_id']==item['class_id'] or other['teacher_id']==item['teacher_id'] or (item['room'] and item['room']==other['room']):
                abort(400,'Conflit de planning avec une séance existante')
    for i,x in enumerate(planned):
        for y in planned[i+1:]:
            if x['start']<y['end'] and x['end']>y['start'] and (x['class_id']==y['class_id'] or x['teacher_id']==y['teacher_id'] or (x['room'] and x['room']==y['room'])):
                abort(400,'Conflit entre les séances sélectionnées')

@app.route('/lessons/<int:lid>/edit',methods=['GET','POST'])
@roles('admin','administratif')
def lesson_edit(lid):
    with db() as c:
        lesson=editable_lesson_row(c,lid)
        if not lesson:abort(404)
        editable=lesson_editable(c,lesson)
    if request.method=='GET':
        return render_template('lesson_edit.html',lesson=lesson,editable=editable,
            duration_hours=(datetime.datetime.fromisoformat(lesson['ends_at'])-datetime.datetime.fromisoformat(lesson['starts_at'])).total_seconds()/3600,
            courses=scoped("""SELECT c.id,c.name,cl.name class_name FROM courses c JOIN classes cl ON cl.id=c.class_id
                    WHERE c.school_id=? ORDER BY cl.name,c.name"""))
    if not editable:abort(400,'Séance historisée : modification interdite')
    try:
        course_id=int(request.form.get('course_id',''))
        date=datetime.date.fromisoformat(request.form.get('lesson_date',''))
        time=datetime.time.fromisoformat(request.form.get('start_time',''))
        duration=float(request.form.get('duration_hours',''))
        if not 0.5<=duration<=12 or not (duration*2).is_integer():raise ValueError()
        start=datetime.datetime.combine(date,time)
        end=start+datetime.timedelta(minutes=round(duration*60))
        if start.date()!=end.date():raise ValueError()
    except (ValueError,OverflowError):abort(400,'Date, heure ou durée invalide')
    course=one('SELECT id,class_id,teacher_id FROM courses WHERE id=? AND school_id=?',(course_id,school_id()))
    if not course or not course['teacher_id']:abort(400,'Cours sans enseignant')
    room=request.form.get('room','').strip()
    if len(room)>80:abort(400)
    item={'id':lid,'start':start.strftime('%Y-%m-%dT%H:%M'),'end':end.strftime('%Y-%m-%dT%H:%M'),
          'room':room,'class_id':course['class_id'],'teacher_id':course['teacher_id']}
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        current=editable_lesson_row(c,lid)
        if not current or not lesson_editable(c,current):abort(400,'Séance non modifiable')
        assert_lesson_schedule_available(c,[item],{lid})
        c.execute('UPDATE lessons SET course_id=?,teacher_id=?,starts_at=?,ends_at=?,room=? WHERE id=? AND school_id=?',
                  (course_id,course['teacher_id'],item['start'],item['end'],room,lid,school_id()))
    notify_event([lesson['course_id'],course_id],
           'Planning modifié : séance du '+lesson['starts_at']+' déplacée au '+item['start']+'.','/planning',teacher_ids=[lesson['teacher_id'],course['teacher_id']])
    audit('lesson.updated',str(lid))
    flash('Séance modifiée. Les personnes concernées ont été notifiées.','success')
    return redirect(url_for('planning'))

@app.post('/lessons/bulk')
@roles('admin','administratif')
def lessons_bulk():
    chosen=request.form.getlist('lesson_ids')
    action=request.form.get('action','')
    if action not in ('delete','modify') or not chosen or len(chosen)>200 or any(not v.isdigit() for v in chosen):abort(400)
    ids=list(dict.fromkeys(int(v) for v in chosen))
    if action=='modify':
        try:
            day_offset=int(request.form.get('day_offset','0'))
            minute_offset=int(request.form.get('minute_offset','0'))
            if abs(day_offset)>365 or abs(minute_offset)>1440 or minute_offset%30:raise ValueError()
        except ValueError:abort(400,'Décalage invalide')
        new_room=request.form.get('room','').strip()
        if len(new_room)>80:abort(400)
        update_room=bool(request.form.get('change_room'))
        if not (day_offset or minute_offset or update_room):abort(400,'Aucune modification sélectionnée')
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        qs=','.join('?' for _ in ids)
        rows=c.execute(f"""SELECT l.*,co.class_id FROM lessons l JOIN courses co ON co.id=l.course_id
                       WHERE l.school_id=? AND l.id IN ({qs})""",(school_id(),*ids)).fetchall()
        if len(rows)!=len(ids):abort(403)
        if any(not lesson_editable(c,row) for row in rows):
            flash('Opération refusée : une séance est réalisée, déclarée, annulée ou possède des présences.','error')
            return redirect(url_for('planning'))
        planned=[]
        if action=='modify':
            for row in rows:
                try:
                    start=datetime.datetime.fromisoformat(row['starts_at'])+datetime.timedelta(days=day_offset,minutes=minute_offset)
                    end=datetime.datetime.fromisoformat(row['ends_at'])+datetime.timedelta(days=day_offset,minutes=minute_offset)
                except (ValueError,OverflowError):abort(400)
                if start.date()!=end.date():abort(400,'Une séance ne peut pas franchir minuit')
                planned.append({'id':row['id'],'start':start.strftime('%Y-%m-%dT%H:%M'),'end':end.strftime('%Y-%m-%dT%H:%M'),
                         'room':new_room if update_room else row['room'] or '', 'class_id':row['class_id'],'teacher_id':row['teacher_id']})
            assert_lesson_schedule_available(c,planned,set(ids))
            for item in planned:
                c.execute('UPDATE lessons SET starts_at=?,ends_at=?,room=? WHERE id=? AND school_id=?',
                          (item['start'],item['end'],item['room'],item['id'],school_id()))
        else:
            c.execute(f'DELETE FROM lessons WHERE school_id=? AND id IN ({qs})',(school_id(),*ids))
    notify_event([row['course_id'] for row in rows],str(len(ids))+' séance(s) '+('supprimée(s)' if action=='delete' else 'modifiée(s)')+' dans le planning.','/planning',teacher_ids=[row['teacher_id'] for row in rows])
    audit('lessons.bulk.'+action,','.join(map(str,ids)))
    flash(f'{len(ids)} séance(s) '+('supprimée(s)' if action=='delete' else 'modifiée(s)')+'.','success')
    return redirect(url_for('planning'))

@app.post('/lessons/<int:lid>/declare')
@roles('enseignant','admin','administratif')
def declare(lid):
    u=current_user();les=one('SELECT * FROM lessons WHERE id=? AND school_id=?',(lid,school_id()))
    if not les or (u['role']=='enseignant' and les['teacher_id']!=u['id']):abort(403)
    # A teacher may confirm a session starting 30 minutes before its planned end, never earlier.
    window_start=datetime.datetime.fromisoformat(les['ends_at'])-datetime.timedelta(minutes=30)
    if datetime.datetime.now()<window_start:abort(400,'Validation disponible 30 minutes avant la fin')
    if les['status'] in ('annulee','reportee'):abort(400,'Séance annulée ou reportée')
    status=request.form.get('status')
    if status not in ('realisee','partielle','non_effectuee'):abort(400)
    mins=int(request.form.get('actual_minutes','0')) if status!='non_effectuee' else 0
    if (status!='non_effectuee' and mins<=0) or mins>24*60:abort(400)
    if status=='non_effectuee' and not request.form.get('comment','').strip():abort(400,'Motif obligatoire')
    execute('UPDATE lessons SET status=?,actual_minutes=?,comment=? WHERE id=?',(status,mins,request.form.get('comment',''),lid))
    if status=='non_effectuee':
        notify_event([les['course_id']],'Séance non effectuée : '+les['starts_at']+' — '+request.form.get('comment','').strip()[:350],'/planning',teacher_ids=[les['teacher_id']])
    elif status=='partielle':
        notify_event([les['course_id']],'Séance partiellement réalisée : '+les['starts_at']+'.','/planning',teacher_ids=[les['teacher_id']])
    audit('lesson.declared',str(lid)+':'+status)
    flash('Séance déclarée.','success');return redirect(url_for('planning'))

@app.post('/lessons/<int:lid>/cancel')
@roles('enseignant','admin','administratif')
def cancel_lesson(lid):
    u=current_user()
    lesson=one('SELECT * FROM lessons WHERE id=? AND school_id=?',(lid,school_id()))
    if not lesson or (u['role']=='enseignant' and lesson['teacher_id']!=u['id']):abort(403)
    reason=request.form.get('reason','').strip()
    if not reason or len(reason)>1500:abort(400,'Motif d’annulation obligatoire')
    if lesson['status'] in ('realisee','partielle'):abort(400,'Impossible d’annuler une séance déjà déclarée')
    execute("UPDATE lessons SET status='annulee',actual_minutes=0,comment=? WHERE id=? AND school_id=?",(reason,lid,school_id()))
    notify_event([lesson['course_id']],'Séance annulée : '+lesson['starts_at']+' — '+reason[:350],'/planning',teacher_ids=[lesson['teacher_id']])
    audit('lesson.cancelled',f'{lid}:{reason}')
    flash('Séance annulée. Les étudiants ont été notifiés.','success')
    return redirect(url_for('planning'))

@app.get('/assessments/<int:aid>/attachment')
@login_required
def assessment_attachment(aid):
    a=one("""SELECT a.attachment_name,a.attachment_blob,a.attachment_mime,c.class_id,c.teacher_id,cl.parents_enabled
           FROM assessments a JOIN courses c ON c.id=a.course_id JOIN classes cl ON cl.id=c.class_id
           WHERE a.id=? AND a.school_id=?""",(aid,school_id()))
    if not a or not a['attachment_blob']:abort(404)
    u=current_user()
    if u['role']=='superadmin':abort(403)
    if u['role']=='enseignant' and a['teacher_id']!=u['id']:abort(403)
    if u['role']=='etudiant' and not one('SELECT id FROM enrollments WHERE school_id=? AND class_id=? AND student_id=?',(school_id(),a['class_id'],u['id'])):abort(403)
    if u['role']=='parent' and (not a['parents_enabled'] or not one("""SELECT p.id FROM parent_links p JOIN enrollments e ON e.student_id=p.student_id WHERE p.school_id=? AND p.parent_id=? AND p.active=1 AND e.class_id=?""",(school_id(),u['id'],a['class_id']))):abort(403)
    response=send_file(io.BytesIO(a['attachment_blob']),download_name=a['attachment_name'],mimetype=a['attachment_mime'],as_attachment=True)
    response.headers['X-Content-Type-Options']='nosniff'
    return response

@app.route('/assessments/<int:aid>/edit',methods=['GET','POST'])
@roles('enseignant','admin','administratif')
def assessment_edit(aid):
    u=current_user()
    item=one("""SELECT a.*,c.teacher_id,c.name course_name,c.class_id FROM assessments a JOIN courses c ON c.id=a.course_id
           WHERE a.id=? AND a.school_id=?""",(aid,school_id()))
    if not item:abort(404)
    if u['role']=='enseignant' and item['teacher_id']!=u['id']:abort(403)
    period=one('SELECT locked FROM periods WHERE school_id=? AND class_id=? ORDER BY id DESC LIMIT 1',(school_id(),item['class_id']))
    if period and period['locked']:abort(403,'Période verrouillée : réouverture nécessaire')
    if request.method=='POST':
        title=request.form.get('title','').strip()
        kind=request.form.get('kind','').strip()
        description=request.form.get('description','').strip()
        due=request.form.get('due_date','').strip() or None
        try:
            coef=float(request.form.get('coefficient',''))
            if not 0<coef<=100:raise ValueError()
            if due:datetime.date.fromisoformat(due)
        except ValueError:abort(400)
        if not title or len(title)>200 or len(description)>10000 or kind not in ('Devoir','Contrôle continu','Examen','Projet'):abort(400)
        execute('UPDATE assessments SET title=?,kind=?,due_date=?,description=?,coefficient=? WHERE id=? AND school_id=?',(title,kind,due,description,coef,aid,school_id()))
        notify_event([item['course_id']],'Évaluation modifiée : '+title+' ('+item['course_name']+').','/assessments',teacher_ids=[item['teacher_id']])
        audit('assessment.updated',str(aid))
        flash('Évaluation modifiée.','success')
        return redirect(url_for('assessments'))
    return render_template('assessment_edit.html',item=item)

@app.get('/assessments/files/<int:fid>')
@login_required
def assessment_file(fid):
    file=one("""SELECT af.filename,af.mime,af.content,a.id aid,c.class_id,c.teacher_id,cl.parents_enabled
                FROM assessment_files af JOIN assessments a ON a.id=af.assessment_id
                JOIN courses c ON c.id=a.course_id JOIN classes cl ON cl.id=c.class_id
                WHERE af.id=? AND a.school_id=?""",(fid,school_id()))
    if not file:abort(404)
    u=current_user()
    if u['role']=='superadmin':abort(403)
    if u['role']=='enseignant' and file['teacher_id']!=u['id']:abort(403)
    if u['role']=='etudiant' and not one('SELECT id FROM enrollments WHERE school_id=? AND class_id=? AND student_id=?',(school_id(),file['class_id'],u['id'])):abort(403)
    if u['role']=='parent' and (not file['parents_enabled'] or not one("""SELECT p.id FROM parent_links p JOIN enrollments e ON e.student_id=p.student_id
         WHERE p.school_id=? AND p.parent_id=? AND p.active=1 AND e.class_id=?""",(school_id(),u['id'],file['class_id']))):abort(403)
    response=send_file(io.BytesIO(file['content']),download_name=file['filename'],mimetype=file['mime'],as_attachment=True)
    response.headers['X-Content-Type-Options']='nosniff'
    return response

@app.route('/assessments',methods=['GET','POST'])
@login_required
def assessments():
    u=current_user()
    if u['role']=='superadmin':abort(403)
    if request.method=='POST':
        if u['role'] not in ('enseignant','admin','administratif'):abort(403)
        course_id=request.form.get('course_id')
        course=one('SELECT * FROM courses WHERE id=? AND school_id=?',(course_id,school_id()))
        if not course or (u['role']=='enseignant' and course['teacher_id']!=u['id']):abort(403)
        title=request.form.get('title','').strip()
        description=request.form.get('description','').strip()
        try:
            coeff=float(request.form.get('coefficient','1'))
            if not (0<coeff<=100):raise ValueError()
        except (ValueError,TypeError):abort(400,'Coefficient invalide')
        if not title or len(title)>200 or len(description)>10000:abort(400)
        uploads=request.files.getlist('documents') or request.files.getlist('document')
        uploads=[file for file in uploads if file and file.filename]
        if len(uploads)>5:abort(400,'Maximum 5 documents')
        prepared=[]
        total_bytes=0
        for file in uploads:
            filename=secure_filename(file.filename)
            data=file.read(3*1024*1024+1)
            total_bytes+=len(data)
            if not filename or len(data)>3*1024*1024 or total_bytes>10*1024*1024:abort(400,'Limites : 3 Mo par document, 10 Mo au total')
            if data.startswith(b'%PDF-'):mime='application/pdf'
            elif data.startswith(bytes([137,80,78,71,13,10,26,10])):mime='image/png'
            elif data.startswith(bytes([255,216,255])):mime='image/jpeg'
            elif data.startswith(b'PK') and filename.lower().endswith('.docx'):mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            else:abort(400,'Formats autorisés : PDF, DOCX, PNG et JPEG')
            prepared.append((filename,mime,data))
        with db() as c:
            aid=c.execute('''INSERT INTO assessments(school_id,course_id,title,kind,due_date,description,coefficient)
                VALUES (?,?,?,?,?,?,?)''',(school_id(),course_id,title,request.form.get('kind','Devoir'),request.form.get('due_date') or None,description,coeff)).lastrowid
            for filename,mime,data in prepared:
                c.execute('INSERT INTO assessment_files(assessment_id,filename,mime,content) VALUES (?,?,?,?)',(aid,filename,mime,data))
        notify_event([course_id],'Nouvelle évaluation : '+title+' ('+course['name']+').','/assessments',teacher_ids=[course['teacher_id']])
        audit('assessment.created',title);flash('Évaluation créée.','success')
        return redirect(url_for('assessments'))
    sql='''SELECT a.id,a.school_id,a.course_id,a.title,a.kind,a.due_date,a.coefficient,a.description,a.attachment_name,c.name course_name,cl.name class_name,c.teacher_id
           FROM assessments a JOIN courses c ON c.id=a.course_id JOIN classes cl ON cl.id=c.class_id WHERE a.school_id=?''';args=[school_id()]
    if u['role']=='enseignant':sql+=' AND c.teacher_id=?';args.append(u['id'])
    elif u['role']=='etudiant':sql+=' AND EXISTS (SELECT 1 FROM enrollments e WHERE e.class_id=c.class_id AND e.student_id=? AND e.school_id=a.school_id)';args.append(u['id'])
    elif u['role']=='parent':sql+=' AND EXISTS (SELECT 1 FROM parent_links p JOIN enrollments e ON e.student_id=p.student_id WHERE p.parent_id=? AND p.active=1 AND e.class_id=c.class_id AND cl.parents_enabled=1)';args.append(u['id'])
    rows=allrows(sql,args)
    scores={}
    if u['role']=='etudiant':
        marks=allrows("""SELECT g.assessment_id,g.value,a.max_score FROM grades g JOIN assessments a ON a.id=g.assessment_id
                         WHERE g.student_id=? AND a.school_id=? AND a.published=1""",(u['id'],school_id()))
        scores={g['assessment_id']:round(g['value']/g['max_score']*20,2) for g in marks if g['max_score']}
    file_rows=allrows("""SELECT af.id,af.assessment_id,af.filename FROM assessment_files af
        JOIN assessments a ON a.id=af.assessment_id WHERE a.school_id=?""",(school_id(),))
    attachments={}
    permitted={row['id'] for row in rows}
    for file in file_rows:
        if file['assessment_id'] in permitted:attachments.setdefault(file['assessment_id'],[]).append(file)
    choices=scoped('SELECT * FROM courses WHERE school_id=?')
    if u['role']=='enseignant':choices=[x for x in choices if x['teacher_id']==u['id']]
    return render_template('assessments.html',assessments=rows,courses=choices,attachments=attachments,scores=scores)

def compile_class_report(class_id, student_ids=None):
    """School-scoped report: evaluation weights are internal to subjects;
    subject coefficients are used only for the overall mean."""
    clazz=one("""SELECT c.id,c.name,c.year_id,y.label academic_year FROM classes c
                 JOIN years y ON y.id=c.year_id WHERE c.id=? AND c.school_id=?""",(class_id,school_id()))
    if not clazz:abort(404)
    courses=allrows('SELECT id,name,code,coefficient FROM courses WHERE class_id=? AND school_id=? ORDER BY name',(class_id,school_id()))
    students=allrows("""SELECT DISTINCT u.id,u.name,u.matricule FROM enrollments e JOIN users u ON u.id=e.student_id
                WHERE e.school_id=? AND e.class_id=? AND e.year_id=? ORDER BY u.name""",(school_id(),class_id,clazz['year_id']))
    if student_ids is not None:students=[student for student in students if student['id'] in student_ids]
    evals=allrows("""SELECT a.id,a.course_id,a.title,a.coefficient,a.max_score,a.published FROM assessments a
                   JOIN courses c ON c.id=a.course_id WHERE a.school_id=? AND c.class_id=? ORDER BY a.id""",(school_id(),class_id))
    marks=allrows("""SELECT g.student_id,g.assessment_id,g.value,g.comment FROM grades g JOIN assessments a ON a.id=g.assessment_id
                  JOIN courses c ON c.id=a.course_id WHERE a.school_id=? AND c.class_id=?""",(school_id(),class_id))
    indexed={(g['student_id'],g['assessment_id']):g for g in marks}
    last_period=one("SELECT id,name,published,locked FROM periods WHERE school_id=? AND class_id=? ORDER BY id DESC LIMIT 1",(school_id(),class_id))
    published=bool(last_period and last_period['published'])
    individual={p['student_id'] for p in allrows('SELECT student_id FROM bulletin_publications WHERE school_id=? AND class_id=?',(school_id(),class_id))}
    results=[]
    for student in students:
        subjects=[];general_num=0.;general_den=0.
        for course in courses:
            assessments=[];num=0.;den=0.
            for ev in evals:
                if ev['course_id']!=course['id'] or not ev['published']:continue
                grade=indexed.get((student['id'],ev['id']))
                score=(grade['value']/ev['max_score']*20) if grade and ev['max_score'] else None
                if score is not None:
                    num+=score*ev['coefficient'];den+=ev['coefficient']
                assessments.append({'id':ev['id'],'title':ev['title'],'coefficient':ev['coefficient'],
                                    'score':round(score,2) if score is not None else None,
                                    'comment':grade['comment'] if grade else None})
            mean=round(num/den,2) if den else None
            if mean is not None:
                general_num+=mean*course['coefficient'];general_den+=course['coefficient']
            subjects.append({'id':course['id'],'title':course['name'],'code':course['code'],
                             'coefficient':course['coefficient'],'assessments':assessments,'mean':mean})
        results.append({'student':student,'subjects':subjects,'general':round(general_num/general_den,2) if general_den else None,'published':published or student['id'] in individual})
    return {'clazz':clazz,'courses':courses,'rows':results,'period':last_period,'published':published,
            'academic_year':clazz['academic_year']}

@app.get('/grades/bulletins')
@roles('admin','administratif')
def school_bulletins():
    options=allrows("""SELECT c.id,c.name,y.id year_id,y.label academic_year
               FROM classes c JOIN years y ON y.id=c.year_id WHERE c.school_id=?
               ORDER BY y.label DESC,c.name""",(school_id(),))
    yr=request.args.get('year_id','')
    shown=[r for r in options if not yr or str(r['year_id'])==yr]
    selected=next((r for r in shown if str(r['id'])==request.args.get('class_id','')),shown[0] if shown else None)
    report=compile_class_report(selected['id']) if selected else None
    years=allrows("SELECT id,label FROM years WHERE school_id=? ORDER BY label DESC",(school_id(),))
    return render_template('school_bulletins.html',report=report,options=shown,years=years,
                           selected=selected,chosen_year=yr)

@app.post('/grades/bulletins/publish')
@roles('admin','administratif')
def school_bulletin_publish():
    class_id=request.form.get('class_id','')
    clazz=one('SELECT id FROM classes WHERE id=? AND school_id=?',(class_id,school_id()))
    if not clazz:abort(403)
    period=one('SELECT id,published FROM periods WHERE class_id=? AND school_id=? ORDER BY id DESC LIMIT 1',(class_id,school_id()))
    if not period:abort(400,'Période à créer avant publication')
    all_student_ids=[r['student_id'] for r in allrows('SELECT DISTINCT student_id FROM enrollments WHERE school_id=? AND class_id=?',(school_id(),class_id))]
    selected=request.form.getlist('student_ids')
    if selected and any(not v.isdigit() for v in selected):abort(400)
    ids=list(dict.fromkeys(map(int,selected))) if selected else all_student_ids
    if not ids or any(uid not in all_student_ids for uid in ids):abort(403)
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        c.execute('UPDATE periods SET locked=1 WHERE id=? AND school_id=?',(period['id'],school_id()))
        for uid in ids:
            c.execute('INSERT OR IGNORE INTO bulletin_publications(school_id,class_id,student_id) VALUES (?,?,?)',(school_id(),int(class_id),uid))
        published_count=c.execute('SELECT COUNT(*) n FROM bulletin_publications WHERE school_id=? AND class_id=?',(school_id(),class_id)).fetchone()['n']
        if published_count>=len(all_student_ids):
            c.execute('UPDATE periods SET published=1 WHERE id=? AND school_id=?',(period['id'],school_id()))
    notify(ids,'Votre bulletin est disponible dans Notes et résultats.','/grades')
    audit('report.published',f'{class_id}:{len(ids)}')
    flash(f'{len(ids)} bulletin(s) publié(s) dans les espaces étudiants. Les notes sont verrouillées.','success')
    return redirect(url_for('school_bulletins',class_id=class_id))

@app.get('/grades/bulletins/print')
@roles('admin','administratif')
def bulletin_print():
    clazz=request.args.get('class_id','')
    if not one('SELECT id FROM classes WHERE id=? AND school_id=?',(clazz,school_id())):abort(404)
    requested=request.args.getlist('student_ids')
    if any(not x.isdigit() for x in requested):abort(400)
    chosen=set(map(int,requested)) if requested else None
    report=compile_class_report(clazz,chosen)
    school=one('SELECT name,logo_mime,id FROM schools WHERE id=?',(school_id(),))
    return render_template('print_bulletins.html',report=report,school=school,printed_on=datetime.date.today().strftime('%d/%m/%Y'))

@app.get('/grades/my-bulletin/<int:cid>')
@login_required
def student_bulletin(cid):
    u=current_user()
    student_id=u['id']
    if u['role']=='parent':
        candidate=request.args.get('student_id','')
        if not candidate.isdigit() or not one("""SELECT p.id FROM parent_links p JOIN enrollments e ON e.student_id=p.student_id
            JOIN classes cl ON cl.id=e.class_id WHERE p.parent_id=? AND p.student_id=? AND p.school_id=? AND p.active=1 AND e.class_id=? AND cl.parents_enabled=1""",(u['id'],int(candidate),school_id(),cid)):abort(403)
        student_id=int(candidate)
    elif u['role']!='etudiant':abort(403)
    if not one('SELECT id FROM enrollments WHERE school_id=? AND student_id=? AND class_id=?',(school_id(),student_id,cid)):abort(403)
    report=compile_class_report(cid,{student_id})
    if not report['rows'] or not report['rows'][0]['published']:abort(403,'Bulletin non encore publié par la scolarité')
    school=one('SELECT name,logo_mime,id FROM schools WHERE id=?',(school_id(),))
    return render_template('print_bulletins.html',report=report,school=school,printed_on=datetime.date.today().strftime('%d/%m/%Y'))

@app.route('/grades',methods=['GET','POST'])
@login_required
def grades():
    u=current_user()
    if u['role']=='superadmin':abort(403)
    if request.method=='POST':
        if u['role'] not in ('enseignant','admin','administratif'):abort(403)
        aid=request.form.get('assessment_id');student=request.form.get('student_id')
        try:value=float(request.form.get('value','-1'))
        except (ValueError,TypeError):abort(400,'Note invalide')
        a=one('''SELECT a.*,c.class_id,c.teacher_id,cl.school_id FROM assessments a JOIN courses c ON c.id=a.course_id JOIN classes cl ON cl.id=c.class_id WHERE a.id=? AND a.school_id=?''',(aid,school_id()))
        if not a or (u['role']=='enseignant' and a['teacher_id']!=u['id']) or not (0<=value<=a['max_score']):abort(403)
        if not one('SELECT id FROM enrollments WHERE student_id=? AND class_id=? AND school_id=?',(student,a['class_id'],school_id())):abort(403)
        period=one('SELECT * FROM periods WHERE class_id=? AND school_id=? ORDER BY id DESC LIMIT 1',(a['class_id'],school_id()))
        if period and period['locked']:
            unlocked=one('SELECT * FROM unlocks WHERE period_id=? AND teacher_id=? AND expires_at>? ORDER BY id DESC LIMIT 1',(period['id'],u['id'],datetime.datetime.now(datetime.timezone.utc).isoformat()))
            if not unlocked:abort(403,'Période verrouillée : réouverture administrative nécessaire')
        existing=one('SELECT * FROM grades WHERE assessment_id=? AND student_id=?',(aid,student))
        if existing:
            execute('UPDATE grades SET value=?,comment=?,updated_at=CURRENT_TIMESTAMP WHERE id=?',(value,request.form.get('comment','')[:2000],existing['id']))
            gid=existing['id']; old=existing['value']
        else:
            gid=execute('INSERT INTO grades(assessment_id,student_id,value,comment) VALUES (?,?,?,?)',(aid,student,value,request.form.get('comment','')[:2000]));old=None
        execute('INSERT INTO grade_history(grade_id,old_value,new_value,actor_id) VALUES (?,?,?,?)',(gid,old,value,u['id']))
        notify([int(student)],'Une note est disponible pour '+a['title']+'.','/grades')
        audit('grade.saved',str(gid));flash('Note enregistrée.','success');return redirect(url_for('grades'))
    if u['role'] in ('etudiant','parent'):
        student_id=u['id']
        if u['role']=='parent':
            # Only an authorized parent with a child in an enabled class may see results.
            children=allrows("""SELECT DISTINCT st.id,st.name,st.matricule FROM parent_links p JOIN users st ON st.id=p.student_id
                       JOIN enrollments e ON e.student_id=st.id JOIN classes cl ON cl.id=e.class_id
                       WHERE p.school_id=? AND p.parent_id=? AND p.active=1 AND cl.parents_enabled=1""",(school_id(),u['id']))
            request_id=request.args.get('student_id','')
            child=next((v for v in children if str(v['id'])==request_id),children[0] if children else None)
            if not child:return render_template('student_grades.html',enrollments=[],selected=None,report=None,children=children,selected_student=None)
            student_id=child['id']
        else:children=[];child=None
        enrolls=allrows("""SELECT DISTINCT cl.id,cl.name,y.label,y.id year_id,cl.parents_enabled
                        FROM enrollments e JOIN classes cl ON cl.id=e.class_id JOIN years y ON y.id=e.year_id
                        WHERE e.school_id=? AND e.student_id=? ORDER BY y.label DESC,cl.name""",(school_id(),student_id))
        if u['role']=='parent':enrolls=[v for v in enrolls if v['parents_enabled']]
        selected=next((v for v in enrolls if str(v['id'])==request.args.get('class_id','')),enrolls[0] if enrolls else None)
        report=compile_class_report(selected['id'],student_ids=[student_id]) if selected else None
        if u['role']=='parent' and report and (not report['rows'] or not report['rows'][0]['published']):
            report=None  # Parents only see officially published results.
        return render_template('student_grades.html',enrollments=enrolls,selected=selected,report=report,
                               children=children,selected_student=child)
    class_rows=allrows("""SELECT cl.id,cl.name,y.label,y.id year_id FROM classes cl JOIN years y ON y.id=cl.year_id
                     WHERE cl.school_id=? ORDER BY y.label DESC,cl.name""",(school_id(),))
    if u['role']=='enseignant':
        class_rows=[cl for cl in class_rows if one('SELECT id FROM courses WHERE school_id=? AND class_id=? AND teacher_id=?',(school_id(),cl['id'],u['id']))]
    selected_class=next((v for v in class_rows if str(v['id'])==request.args.get('class_id','')),class_rows[0] if class_rows else None)
    courses=allrows('SELECT id,name FROM courses WHERE school_id=? AND class_id=? ORDER BY name',(school_id(),selected_class['id'])) if selected_class else []
    if u['role']=='enseignant':courses=[c for c in courses if one('SELECT id FROM courses WHERE id=? AND teacher_id=?',(c['id'],u['id']))]
    selected_course=next((c for c in courses if str(c['id'])==request.args.get('course_id','')),courses[0] if courses else None)
    assessments=allrows('SELECT id,title,coefficient FROM assessments WHERE school_id=? AND course_id=? ORDER BY id',(school_id(),selected_course['id'])) if selected_course else []
    selected_assessment=next((v for v in assessments if str(v['id'])==request.args.get('assessment_id','')),assessments[0] if assessments else None)
    return render_template('staff_grades.html',classes=class_rows,selected_class=selected_class,courses=courses,
                           selected_course=selected_course,assessments=assessments,selected_assessment=selected_assessment)
    
@app.route('/grades/entry/<int:aid>',methods=['GET','POST'])
@roles('enseignant','admin','administratif')
def grade_entry(aid):
    a=one("""SELECT a.id,a.title,a.max_score,a.coefficient,c.id course_id,c.class_id,c.teacher_id,c.name course_name,cl.name class_name
             FROM assessments a JOIN courses c ON c.id=a.course_id JOIN classes cl ON cl.id=c.class_id
             WHERE a.id=? AND a.school_id=?""",(aid,school_id()))
    if not a:abort(404)
    u=current_user()
    if u['role']=='enseignant' and a['teacher_id']!=u['id']:abort(403)
    period=one('SELECT * FROM periods WHERE class_id=? AND school_id=? ORDER BY id DESC LIMIT 1',(a['class_id'],school_id()))
    if period and period['locked']:
        if not one('SELECT id FROM unlocks WHERE period_id=? AND teacher_id=? AND expires_at>?',(period['id'],u['id'],datetime.datetime.now(datetime.timezone.utc).isoformat())):abort(403,'Période verrouillée : demander une réouverture')
    students=allrows("""SELECT DISTINCT u.id,u.name,u.matricule,g.value,g.comment FROM enrollments e JOIN users u ON u.id=e.student_id
                        LEFT JOIN grades g ON g.student_id=u.id AND g.assessment_id=?
                        WHERE e.school_id=? AND e.class_id=? ORDER BY u.name""",(aid,school_id(),a['class_id']))
    if request.method=='POST':
        changes=[]
        for student in students:
            raw=request.form.get(f"note_{student['id']}",'').strip()
            comment=request.form.get(f"comment_{student['id']}",'').strip()
            if len(comment)>2000:abort(400,'Commentaire trop long')
            if not raw:
                if student['value'] is not None: # Empty input does not erase an existing grade
                    continue
                continue
            try:
                value=float(raw)
                if not (0<=value<=a['max_score']):raise ValueError()
            except ValueError:abort(400,'Note invalide')
            changes.append((student['id'],value,comment,student['value']))
        with db() as c:
            for student_id,value,comment,old in changes:
                row=c.execute('SELECT id FROM grades WHERE assessment_id=? AND student_id=?',(aid,student_id)).fetchone()
                if row:
                    c.execute('UPDATE grades SET value=?,comment=?,updated_at=CURRENT_TIMESTAMP WHERE id=?',(value,comment,row['id']))
                    gid=row['id']
                else:
                    gid=c.execute('INSERT INTO grades(assessment_id,student_id,value,comment) VALUES (?,?,?,?)',(aid,student_id,value,comment)).lastrowid
                if old!=value:
                    c.execute('INSERT INTO grade_history(grade_id,old_value,new_value,actor_id) VALUES (?,?,?,?)',(gid,old,value,u['id']))
        if changes and a:
            notify([x[0] for x in changes],'Une note a été enregistrée pour '+a['title']+'. Consultez Notes et résultats.','/grades')
        audit('grades.batch.updated',f'{aid}:{len(changes)}')
        flash(f'{len(changes)} note(s) enregistrée(s).','success')
        return redirect(url_for('grade_entry',aid=aid))
    return render_template('grade_entry.html',assessment=a,students=students)

@app.get('/grades/global')
@roles('enseignant','admin','administratif')
def grade_global():
    u=current_user()
    choices=scoped("""SELECT c.id,c.name,cl.name class_name,c.coefficient FROM courses c JOIN classes cl ON cl.id=c.class_id
                      WHERE c.school_id=? ORDER BY cl.name,c.name""")
    if u['role']=='enseignant':choices=[c for c in choices if one('SELECT id FROM courses WHERE id=? AND teacher_id=?',(c['id'],u['id']))]
    requested=request.args.get('course_id','')
    selected=next((c for c in choices if str(c['id'])==requested),choices[0] if choices else None)
    if not selected:return render_template('grade_global.html',choices=choices,course=None,assessments=[],rows=[])
    course=one('SELECT * FROM courses WHERE id=? AND school_id=?',(selected['id'],school_id()))
    assessments=allrows('SELECT id,title,coefficient,max_score FROM assessments WHERE school_id=? AND course_id=? ORDER BY id',(school_id(),course['id']))
    students=allrows("""SELECT DISTINCT u.id,u.name,u.matricule FROM enrollments e JOIN users u ON u.id=e.student_id
                        WHERE e.school_id=? AND e.class_id=? ORDER BY u.name""",(school_id(),course['class_id']))
    results=allrows("""SELECT g.student_id,g.assessment_id,g.value FROM grades g JOIN assessments a ON a.id=g.assessment_id
                       WHERE a.school_id=? AND a.course_id=?""",(school_id(),course['id']))
    marks={(g['student_id'],g['assessment_id']):g['value'] for g in results}
    rows=[]
    for student in students:
        cells=[marks.get((student['id'],a['id'])) for a in assessments]
        weighted=sum((v/a['max_score']*20)*a['coefficient'] for v,a in zip(cells,assessments) if v is not None and a['max_score'])
        weight=sum(a['coefficient'] for v,a in zip(cells,assessments) if v is not None and a['max_score'])
        rows.append({'student':student,'notes':cells,'average':round(weighted/weight,2) if weight else None})
    return render_template('grade_global.html',choices=choices,course=course,assessments=assessments,rows=rows)

@app.get('/grades/class-summary')
@roles('enseignant','admin','administratif')
def class_grade_summary():
    u=current_user()
    classes=scoped('SELECT id,name FROM classes WHERE school_id=? ORDER BY name')
    if u['role']=='enseignant':
        classes=[cl for cl in classes if one('SELECT id FROM courses WHERE class_id=? AND school_id=? AND teacher_id=?',(cl['id'],school_id(),u['id']))]
    selected_id=request.args.get('class_id','')
    selected=next((cl for cl in classes if str(cl['id'])==selected_id),classes[0] if classes else None)
    if not selected:return render_template('class_grade_summary.html',classes=classes,selected=None,courses=[],rows=[])
    courses=allrows('SELECT id,name,coefficient FROM courses WHERE class_id=? AND school_id=? ORDER BY name',(selected['id'],school_id()))
    if u['role']=='enseignant':
        courses=[c for c in courses if one('SELECT id FROM courses WHERE id=? AND teacher_id=?',(c['id'],u['id']))]
    students=allrows("""SELECT DISTINCT u.id,u.name,u.matricule FROM enrollments e JOIN users u ON u.id=e.student_id
                        WHERE e.school_id=? AND e.class_id=? ORDER BY u.name""",(school_id(),selected['id']))
    course_averages={}
    for course in courses:
        assessments=allrows('SELECT id,coefficient,max_score FROM assessments WHERE school_id=? AND course_id=?',(school_id(),course['id']))
        marks=allrows('SELECT g.assessment_id,g.student_id,g.value FROM grades g JOIN assessments a ON a.id=g.assessment_id WHERE a.school_id=? AND a.course_id=?',(school_id(),course['id']))
        marks_by={(g['student_id'],g['assessment_id']):g['value'] for g in marks}
        for student in students:
            numerator=denominator=0.0
            for a in assessments:
                value=marks_by.get((student['id'],a['id']))
                if value is not None and a['max_score']:
                    numerator+=value/a['max_score']*20*a['coefficient']
                    denominator+=a['coefficient']
            course_averages[(student['id'],course['id'])]=round(numerator/denominator,2) if denominator else None
    rows=[]
    for student in students:
        notes=[course_averages[(student['id'],course['id'])] for course in courses]
        weighted=sum(v*course['coefficient'] for v,course in zip(notes,courses) if v is not None)
        denominator=sum(course['coefficient'] for v,course in zip(notes,courses) if v is not None)
        rows.append({'student':student,'notes':notes,'general':round(weighted/denominator,2) if denominator else None})
    return render_template('class_grade_summary.html',classes=classes,selected=selected,courses=courses,rows=rows)

@app.route('/periods',methods=['GET','POST'])
@roles('admin','administratif')
def periods():
    if request.method=='POST':
        pid=request.form.get('period_id');action=request.form.get('action')
        if not one('SELECT id FROM periods WHERE id=? AND school_id=?',(pid,school_id())):abort(403)
        if action=='lock':execute('UPDATE periods SET locked=1 WHERE id=?',(pid,))
        elif action=='publish':execute('UPDATE periods SET locked=1,published=1 WHERE id=?',(pid,))
        elif action=='unpublish':execute('UPDATE periods SET published=0 WHERE id=?',(pid,))
        elif action=='unlock':execute('UPDATE periods SET locked=0 WHERE id=?',(pid,))
        else:abort(400)
        audit('period.'+action,str(pid));flash('Période mise à jour.','success');return redirect(url_for('periods'))
    return render_template('periods.html',periods=scoped('SELECT p.*,c.name class_name FROM periods p JOIN classes c ON c.id=p.class_id WHERE p.school_id=?'))

@app.get('/api/notifications')
@login_required
def notifications_api():
    uid=current_user()['id']
    rows=allrows('SELECT id,message,href,read,created_at FROM notifications WHERE user_id=? ORDER BY id DESC LIMIT 60',(uid,))
    count=one('SELECT count(*) total FROM notifications WHERE user_id=? AND read=0',(uid,))['total']
    return jsonify(notifications=[dict(row) for row in rows],unread=count)

@app.post('/api/notifications/read')
@login_required
def notifications_read():
    uid=current_user()['id']
    ids=(request.get_json(silent=True) or {}).get('ids',[])
    if not isinstance(ids,list) or len(ids)>100 or any(not isinstance(i,int) for i in ids):abort(400)
    with db() as c:
        if ids:
            c.execute('UPDATE notifications SET read=1 WHERE user_id=? AND id IN ('+','.join('?' for _ in ids)+')',(uid,*ids))
        else:
            c.execute('UPDATE notifications SET read=1 WHERE user_id=?',(uid,))
    return jsonify(ok=True)

@app.post('/api/notifications/delete')
@login_required
def notifications_delete():
    uid=current_user()['id']
    payload=request.get_json(silent=True) or {}
    if payload.get('all') is True:
        execute('DELETE FROM notifications WHERE user_id=?',(uid,))
    else:
        ids=payload.get('ids',[])
        if not isinstance(ids,list) or not ids or len(ids)>100 or any(not isinstance(i,int) for i in ids):abort(400)
        with db() as c:
            c.execute('DELETE FROM notifications WHERE user_id=? AND id IN ('+','.join('?' for _ in ids)+')',(uid,*ids))
    return jsonify(ok=True)

@app.get('/api/push/config')
@login_required
def push_config():
    key=os.getenv('VAPID_PUBLIC_KEY','').strip()
    return jsonify(enabled=bool(key and os.getenv('VAPID_PRIVATE_KEY') and os.getenv('VAPID_SUBJECT')),public_key=key)

@app.post('/api/push/subscribe')
@login_required
def push_subscribe():
    sub=request.get_json(silent=True)
    if not isinstance(sub,dict):abort(400)
    endpoint=sub.get('endpoint','')
    keys=sub.get('keys',{})
    if not isinstance(endpoint,str) or len(endpoint)>2048 or not isinstance(keys,dict):abort(400)
    parsed=urlsplit(endpoint)
    # Restrict callback endpoints to known Web Push operators to prevent arbitrary server requests.
    trusted_hosts=('fcm.googleapis.com','updates.push.services.mozilla.com',
                   'web.push.apple.com','push.apple.com','notify.windows.com','wns.windows.com')
    host=(parsed.hostname or '').lower()
    is_provider=any(host==domain or host.endswith('.'+domain) for domain in trusted_hosts)
    if (parsed.scheme!='https' or not is_provider or parsed.username or parsed.password
        or not all(isinstance(keys.get(k),str) and 8<=len(keys[k])<=512 for k in ('p256dh','auth'))):
        abort(400,'Abonnement push invalide')
    raw=json.dumps({'endpoint':endpoint,'keys':{'p256dh':keys['p256dh'],'auth':keys['auth']}},separators=(',',':'))
    with db() as c:
        c.execute('INSERT INTO push_subscriptions(user_id,endpoint,subscription_json) VALUES (?,?,?) ON CONFLICT(endpoint) DO UPDATE SET user_id=excluded.user_id,subscription_json=excluded.subscription_json',(current_user()['id'],endpoint,raw))
    return jsonify(ok=True)

@app.post('/api/push/unsubscribe')
@login_required
def push_unsubscribe():
    endpoint=(request.get_json(silent=True) or {}).get('endpoint','')
    if not isinstance(endpoint,str) or len(endpoint)>2048:abort(400)
    execute('DELETE FROM push_subscriptions WHERE user_id=? AND endpoint=?',(current_user()['id'],endpoint))
    return jsonify(ok=True)

@app.get('/installer')
def install_app():
    return render_template('install_app.html')


@app.get('/sw.js')
def service_worker():
    return send_file(ROOT/'static'/'sw.js',mimetype='application/javascript',max_age=0)

@app.get('/manifest.webmanifest')
def webmanifest():
    return send_file(ROOT/'static'/'manifest.webmanifest',mimetype='application/manifest+json')

@app.get('/reports')
@roles('admin','administratif')
def reports():
    years=scoped('SELECT id,label FROM years WHERE school_id=? ORDER BY label DESC')
    classes=scoped('SELECT cl.id,cl.name,cl.year_id FROM classes cl WHERE cl.school_id=? ORDER BY cl.name')
    teachers=scoped("SELECT id,name FROM users WHERE school_id=? AND role IN ('enseignant','admin') ORDER BY name")
    filters={}
    for key,allowed in [('year_id',years),('class_id',classes),('teacher_id',teachers)]:
        raw=request.args.get(key,'').strip()
        if raw:
            if not raw.isdigit() or int(raw) not in {r['id'] for r in allowed}:abort(400,'Filtre invalide')
            filters[key]=int(raw)
    conditions=['c.school_id=?']
    params=[school_id()]
    if 'year_id' in filters:
        conditions.append('cl.year_id=?')
        params.append(filters['year_id'])
    if 'class_id' in filters:
        conditions.append('cl.id=?')
        params.append(filters['class_id'])
    if 'teacher_id' in filters:
        conditions.append('c.teacher_id=?')
        params.append(filters['teacher_id'])
    rows=allrows('''SELECT c.name,cl.name class_name,c.planned_minutes,COUNT(l.id) lessons,
      COALESCE(SUM(CASE WHEN l.status IN ('realisee','partielle') THEN l.actual_minutes ELSE 0 END),0) actual_minutes
      FROM courses c JOIN classes cl ON cl.id=c.class_id AND cl.school_id=c.school_id
      LEFT JOIN lessons l ON l.course_id=c.id AND l.school_id=c.school_id
      WHERE '''+' AND '.join(conditions)+''' GROUP BY c.id ORDER BY cl.name,c.name''',tuple(params))
    if request.args.get('format')=='csv':
        import csv,io
        b=io.StringIO();w=csv.writer(b,delimiter=';');w.writerow(['Cours','Classe','Séances planifiées','Heures prévues','Heures réalisées','Taux %'])
        for x in rows:w.writerow([x['name'],x['class_name'],x['lessons'],x['planned_minutes']/60,x['actual_minutes']/60,round(x['actual_minutes']/x['planned_minutes']*100,1) if x['planned_minutes'] else 0])
        return Response('\ufeff'+b.getvalue(),headers={'Content-Disposition':'attachment; filename=edupilot-rapport.csv'},mimetype='text/csv')
    return render_template('reports.html',rows=rows,years=years,classes=classes,teachers=teachers,filters=filters)

@app.route('/campuses',methods=['GET','POST'])
@roles('admin','administratif')
def campuses():
    school=one('SELECT * FROM schools WHERE id=?',(school_id(),))
    if not school or not school['campus_enabled']:abort(404)
    if request.method=='POST':
        name=request.form.get('name','').strip()
        if not name:abort(400)
        execute('INSERT INTO campuses(school_id,name) VALUES (?,?)',(school_id(),name))
        audit('campus.created',name)
        return redirect(url_for('campuses'))
    return render_template('list.html',title='Campus et sites',columns=['Campus'],rows=[(x['name'],) for x in scoped('SELECT name FROM campuses WHERE school_id=?')],form='campuses')

@app.post('/courses/<int:cid>/assign')
@roles('admin','administratif')
def assign_teacher(cid):
    if not has_school('courses',cid):abort(403)
    teacher=request.form.get('teacher_id')
    if not one("SELECT id FROM users WHERE school_id=? AND id=? AND role IN ('enseignant','admin')",(school_id(),teacher)):abort(403)
    old=one('SELECT teacher_id FROM courses WHERE id=?',(cid,))
    now=datetime.datetime.now(datetime.timezone.utc).isoformat()
    if old['teacher_id']==int(teacher):
        flash('Cet enseignant est déjà affecté à ce cours.','success')
        return redirect(url_for('courses'))
    try:
        with db() as c:
            c.execute('BEGIN IMMEDIATE')
            sync_course_future_teacher(c,cid,int(teacher))
            if old['teacher_id']:
                c.execute('UPDATE teacher_history SET date_end=? WHERE course_id=? AND date_end IS NULL',(now,cid))
            c.execute('UPDATE courses SET teacher_id=? WHERE id=?',(teacher,cid))
            c.execute('INSERT INTO teacher_history(course_id,teacher_id,date_start) VALUES (?,?,?)',(cid,teacher,now))
    except ValueError as exc:
        flash(str(exc),'error')
        return redirect(url_for('courses'))
    audit('course.teacher.changed',str(cid))
    notify([int(teacher)],'Vous êtes affecté à un nouveau cours.','/planning')
    flash('Enseignant affecté. Les séances futures sont synchronisées et l’historique est conservé.','success')
    return redirect(url_for('courses'))

@app.post('/parents/link')
@roles('admin','administratif')
def parent_link():
    student=request.form.get('student_id');parent=request.form.get('parent_id')
    if not one("SELECT id FROM users WHERE school_id=? AND id=? AND role='etudiant'",(school_id(),student)):abort(403)
    if not one("SELECT id FROM users WHERE school_id=? AND id=? AND role='parent'",(school_id(),parent)):abort(403)
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        existing=c.execute('SELECT id FROM parent_links WHERE school_id=? AND student_id=? AND parent_id=?',(school_id(),student,parent)).fetchone()
        count=c.execute('SELECT count(*) n FROM parent_links WHERE school_id=? AND student_id=? AND active=1',(school_id(),student)).fetchone()['n']
        if count>=2 and not existing:abort(400,'Limite de 2 responsables')
        if existing:c.execute('UPDATE parent_links SET active=1 WHERE id=?',(existing['id'],))
        else:c.execute('INSERT INTO parent_links(school_id,student_id,parent_id) VALUES (?,?,?)',(school_id(),student,parent))
    audit('parent.linked',str(student)+':'+str(parent))
    notify([int(parent)],'Vous êtes rattaché à un apprenant dans EduPilot.','/grades')
    flash('Responsable associé.','success')
    return redirect(url_for('parents'))

@app.route('/parents')
@roles('admin','administratif')
def parents():
    linked=scoped('SELECT s.name student_name,s.matricule student_matricule,p.name parent_name,p.matricule parent_matricule,pl.relation FROM parent_links pl JOIN users s ON s.id=pl.student_id JOIN users p ON p.id=pl.parent_id WHERE pl.school_id=? AND pl.active=1')
    return render_template('parents.html',linked=linked,students=scoped("SELECT id,name,matricule FROM users WHERE school_id=? AND role='etudiant'"),parents=scoped("SELECT id,name,matricule FROM users WHERE school_id=? AND role='parent'"))

@app.get('/attendance/<int:lid>')
@roles('enseignant','admin','administratif')
def attendance(lid):
    lesson=one('SELECT l.*,c.class_id,c.name course_name FROM lessons l JOIN courses c ON c.id=l.course_id WHERE l.id=? AND l.school_id=?',(lid,school_id()))
    if not lesson or (current_user()['role']=='enseignant' and lesson['teacher_id']!=current_user()['id']):abort(403)
    students=allrows('SELECT u.id,u.name,u.matricule,a.status FROM enrollments e JOIN users u ON u.id=e.student_id LEFT JOIN attendance a ON a.lesson_id=? AND a.student_id=u.id WHERE e.school_id=? AND e.class_id=?',(lid,school_id(),lesson['class_id']))
    return render_template('attendance.html',lesson=lesson,students=students)

@app.post('/attendance/<int:lid>')
@roles('enseignant','admin','administratif')
def attendance_save(lid):
    lesson=one('SELECT l.*,c.class_id FROM lessons l JOIN courses c ON c.id=l.course_id WHERE l.id=? AND l.school_id=?',(lid,school_id()))
    if not lesson or (current_user()['role']=='enseignant' and lesson['teacher_id']!=current_user()['id']):abort(403)
    students=allrows('SELECT student_id FROM enrollments WHERE school_id=? AND class_id=?',(school_id(),lesson['class_id']))
    with db() as c:
        for x in students:
            v=request.form.get('student_'+str(x['student_id']))
            if v in ('present','absent','retard','excuse'):
                c.execute('INSERT INTO attendance(lesson_id,student_id,status) VALUES (?,?,?) ON CONFLICT(lesson_id,student_id) DO UPDATE SET status=excluded.status',(lid,x['student_id'],v))
    audit('attendance.saved',str(lid));flash('Appel enregistré.','success')
    return redirect(url_for('attendance',lid=lid))

@app.get('/api/planning')
@login_required
def planning_json():
    u=current_user()
    if u['role']=='superadmin':abort(403)
    sql='''SELECT l.id,l.starts_at,l.ends_at,l.status,l.room,c.name course,c.class_id,cl.name class_name,l.teacher_id,t.name teacher_name FROM lessons l JOIN courses c ON c.id=l.course_id JOIN classes cl ON cl.id=c.class_id LEFT JOIN users t ON t.id=l.teacher_id WHERE l.school_id=?'''; args=[school_id()]
    if u['role']=='enseignant' or (u['role']=='admin' and request.args.get('mine')=='1'):
        sql+=' AND l.teacher_id=?';args.append(u['id'])
    elif u['role']=='etudiant':sql+=' AND EXISTS (SELECT 1 FROM enrollments e WHERE e.student_id=? AND e.class_id=c.class_id)';args.append(u['id'])
    elif u['role']=='parent':sql+=' AND EXISTS (SELECT 1 FROM parent_links p JOIN enrollments e ON e.student_id=p.student_id WHERE p.parent_id=? AND e.class_id=c.class_id AND p.active=1 AND cl.parents_enabled=1)';args.append(u['id'])
    return jsonify([{'id':x['id'],'title':x['course']+' — '+x['class_name']+' · '+(x['teacher_name'] or 'Enseignant à préciser')+' · '+(x['room'] or 'Salle à préciser'),'start':x['starts_at'],'end':x['ends_at'],'extendedProps':{'status':x['status'],'room':x['room'],'teacher':x['teacher_name'],'classId':x['class_id'],'className':x['class_name'],'teacherId':x['teacher_id'],'actionable':(u['role'] in ('enseignant','admin','administratif') and (u['role']!='enseignant' or x['teacher_id']==u['id']))}} for x in allrows(sql,args)])

@app.errorhandler(403)
def denied(e):return render_template('error.html',message='Accès non autorisé.'),403
@app.errorhandler(404)
def missing(e):return render_template('error.html',message='Page introuvable.'),404
@app.errorhandler(400)
def invalid(e):return render_template('error.html',message='Données ou demande invalides.'),400

if __name__=='__main__':app.run(debug=False,port=int(os.getenv('PORT',5000)))
